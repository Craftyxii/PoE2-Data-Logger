"""Workbook/CSV checks for sheet relationships, compact Atlas setup rows, numeric item counts and snapshot-consistent exports."""

from contextlib import contextmanager
import csv
import importlib.util
import io
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PoE2_Data_Logger.core import logger_store as logger, store, workbook_export


ATLAS_HEADERS = ["Atlas Setup ID", "Map IDs", "Atlas Data Version", "Gear Item Rarity %", "Atlas Catalog ID",
                 "Atlas Node: Main Atlas | Ordinary Node [0007]",
                 "Atlas Node: Expedition | Choice Node [0008]",
                 "Atlas Choice: Expedition | Choice Node [0008]"]


def csv_bytes(headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


@contextmanager
def memory_database():
    database = sqlite3.connect(":memory:")
    try:
        yield database
    finally:
        database.close()


class AtlasWorkbookTests(unittest.TestCase):
    def sample_workbook(self):
        main_data = csv_bytes(["Map ID", "Atlas Setup ID", "Gear Item Rarity %", "Chaos Orb"],
                              [["M0001", "ATLAS-1", 0, 7], ["M0002", "ATLAS-2", 175.5, 0],
                               ["M0000", "", "", ""]])
        history_data = csv_bytes(["Scan Commit #", "Atlas Setup ID", "Type", "Currency: Chaos Orb"],
                                 [[1, "ATLAS-1", "Currency", 3], [2, "ATLAS-1", "Currency", 7],
                                  [3, "ATLAS-2", "Currency", 0]])
        atlas_data = csv_bytes(ATLAS_HEADERS, [
            ["ATLAS-1", "M0001, M0003", "patch-one", 0, "000123", "No", "Yes", 1],
            ["ATLAS-2", "M0002", "patch-one", 175.5, "000123", "Yes", "No", 2],
            ["ATLAS-unknown", "", "patch-two", "", "000456", "No", "", ""],
        ])
        connections = []

        def main_export(*, _db):
            self.assertTrue(_db.in_transaction)
            connections.append(_db)
            return main_data

        def atlas_export(*, _db):
            self.assertTrue(_db.in_transaction)
            connections.append(_db)
            return atlas_data

        def history_export(*, _db):
            self.assertTrue(_db.in_transaction)
            connections.append(_db)
            return history_data

        def numeric_item_headers(*, _db):
            self.assertTrue(_db.in_transaction)
            connections.append(_db)
            return ["Chaos Orb"]

        with patch.object(logger, "_connect", memory_database), \
                patch.object(logger, "export_primary_csv", side_effect=main_export), \
                patch.object(logger, "export_atlas_csv", side_effect=atlas_export), \
                patch.object(logger, "export_all_csv", side_effect=history_export), \
                patch.object(logger, "map_summary_item_headers", side_effect=numeric_item_headers):
            workbook = workbook_export.export_xlsx()
        self.assertEqual(len(connections), 4)
        self.assertTrue(all(connection is connections[0] for connection in connections))
        return workbook

    def test_companion_sheet_package_and_relationships(self):
        ns = {"s": workbook_export.NS, "r": workbook_export.DOC_REL,
              "p": workbook_export.PKG_REL, "c": workbook_export.TYPES}
        with ZipFile(io.BytesIO(self.sample_workbook())) as archive:
            self.assertIsNone(archive.testzip())
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            sheets = workbook.findall("s:sheets/s:sheet", ns)
            self.assertEqual([sheet.get("name") for sheet in sheets],
                             ["Export", "Atlas Character Settings", "Scan History"])
            relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {row.get("Id"): row.get("Target") for row in relationships}
            for number, sheet in enumerate(sheets, 1):
                self.assertEqual(targets[sheet.get(f"{{{workbook_export.DOC_REL}}}id")],
                                 f"worksheets/sheet{number}.xml")
            types = ET.fromstring(archive.read("[Content_Types].xml"))
            overrides = {row.get("PartName"): row.get("ContentType") for row in types}
            for number in (1, 2, 3):
                self.assertTrue(overrides[f"/xl/worksheets/sheet{number}.xml"].endswith("worksheet+xml"))

            main = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            links = {link.get("ref"): link.get("location")
                     for link in main.findall("s:hyperlinks/s:hyperlink", ns)}
            self.assertEqual(links, {"B2": "'Atlas Character Settings'!A2",
                                     "B3": "'Atlas Character Settings'!A3"})
            linked_cell = main.find("s:sheetData/s:row/s:c[@r='B2']", ns)
            self.assertEqual(linked_cell.get("s"), "2")
            self.assertNotIn("B4", links)
            self.assertFalse(any(link.get(f"{{{workbook_export.DOC_REL}}}id")
                                 for link in main.findall("s:hyperlinks/s:hyperlink", ns)))

            history = ET.fromstring(archive.read("xl/worksheets/sheet3.xml"))
            history_links = {link.get("ref"): link.get("location")
                             for link in history.findall("s:hyperlinks/s:hyperlink", ns)}
            self.assertEqual(history_links, {"B2": "'Atlas Character Settings'!A2",
                                             "B3": "'Atlas Character Settings'!A2",
                                             "B4": "'Atlas Character Settings'!A3"})
            count = main.find("s:sheetData/s:row/s:c[@r='D2']", ns)
            self.assertIsNone(count.get("t"))
            self.assertEqual(count.findtext("s:v", namespaces=ns), "7")

            atlas = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
            self.assertEqual(atlas.find("s:dimension", ns).get("ref"), "A1:H4")
            self.assertEqual(atlas.find("s:autoFilter", ns).get("ref"), "A1:H4")
            self.assertEqual(atlas.find("s:sheetViews/s:sheetView/s:pane", ns).get("state"), "frozen")
            rows = atlas.findall("s:sheetData/s:row", ns)
            self.assertEqual(["".join(cell.itertext()) for cell in rows[0]], ATLAS_HEADERS)
            cells = {cell.get("r"): cell for row in rows for cell in row}
            for address, value in (("D2", "0"), ("H2", "1"), ("H3", "2"), ("D3", "175.5")):
                self.assertIsNone(cells[address].get("t"))
                self.assertEqual(cells[address].findtext("s:v", namespaces=ns), value)
            self.assertNotIn("D4", cells)
            self.assertNotIn("G4", cells)
            self.assertNotIn("H4", cells)
            self.assertEqual("".join(cells["E2"].itertext()), "000123")
            self.assertEqual("".join(cells["F2"].itertext()), "No")
            self.assertEqual("".join(cells["G3"].itertext()), "No")
            self.assertEqual("".join(cells["B2"].itertext()), "M0001, M0003")

    def test_setup_hyperlinks_follow_reordered_columns_and_escape_sheet_name(self):
        main_data = csv_bytes(["Map ID", "Detail", "Atlas Setup ID"],
                              [["M1", "first", "setup-two"], ["M2", "second", "setup-one"],
                               ["M3", "legacy", ""], ["M4", "unmatched", "missing"]])
        atlas_data = csv_bytes(["Atlas Node: Main Atlas | First [n1]", "Atlas Setup ID", "Map IDs"],
                               [["No", "setup-one", "M2"], ["Yes", "setup-two", "M1"]])
        links = workbook_export._atlas_setup_hyperlinks(main_data, atlas_data, "Atlas's Settings")
        self.assertEqual(links, {"C2": "'Atlas''s Settings'!B3", "C3": "'Atlas''s Settings'!B2"})
        root = ET.fromstring(workbook_export._data_sheet(main_data, hyperlinks=links))
        ns = {"s": workbook_export.NS}
        self.assertEqual({link.get("ref"): link.get("location")
                          for link in root.findall("s:hyperlinks/s:hyperlink", ns)}, links)

    def test_actual_export_links_each_map_to_its_matching_saved_setup(self):
        from PoE2_Data_Logger.core import atlas_catalog
        with tempfile.TemporaryDirectory(prefix="atlas-workbook-integration-") as directory:
            previous = store.DATA_DIR
            try:
                store.DATA_DIR = Path(directory)
                logger._READY = False
                logger.initialize()
                logger.start_map()
                catalog = atlas_catalog.catalog()
                nodes = [key for key, node in catalog["nodes"].items()
                         if node.get("allocatable") and not node.get("choices")]
                choice_id, choice_node = next((key, node) for key, node in catalog["nodes"].items()
                                             if node.get("allocatable") and len(node.get("choices", [])) >= 2)
                logger.save_atlas_settings({"catalog_version": catalog["version"],
                    "allocated": [*nodes[:1], choice_id],
                    "choices": {}, "gear_item_rarity": 0})
                logger.commit_chain("Rage")
                logger.finish_map(0, 0, 0, 0)
                logger.save_atlas_settings({"catalog_version": catalog["version"],
                    "allocated": nodes[:2],
                    "choices": {choice_id: choice_node["choices"][1]["id"]}, "gear_item_rarity": 57.5})
                logger.start_map()
                logger.commit_chain("Time")
                main_rows = list(csv.DictReader(io.StringIO(logger.export_primary_csv().decode("utf-8-sig"))))
                atlas_rows = list(csv.DictReader(io.StringIO(logger.export_atlas_csv().decode("utf-8-sig"))))
                setup_by_map = {row["Map ID"]: row["Atlas Setup ID"] for row in main_rows}
                self.assertEqual(set(setup_by_map), {"M0001", "M0002"})
                self.assertNotEqual(setup_by_map["M0001"], setup_by_map["M0002"])
                self.assertEqual(len(atlas_rows), 2)
                for row in atlas_rows:
                    self.assertEqual(row["Atlas Setup ID"], setup_by_map[row["Map IDs"]])
                by_map = {row["Map IDs"]: row for row in atlas_rows}
                choice_suffix = f"{choice_node['activity']} | {choice_node['name']} [{choice_id}]"
                choice_column = f"Atlas Choice: {choice_suffix}"
                choice_allocation_column = f"Atlas Node: {choice_suffix}"
                self.assertEqual(by_map["M0001"][choice_column], "0")
                self.assertEqual(by_map["M0002"][choice_column], "2")
                self.assertEqual(by_map["M0001"][choice_allocation_column], "Yes")
                self.assertEqual(by_map["M0002"][choice_allocation_column], "No")
                first_node = catalog["nodes"][nodes[0]]
                always_checked_column = (f"Atlas Node: {first_node['activity']} | "
                                         f"{first_node['name']} [{nodes[0]}]")
                self.assertNotIn(always_checked_column, atlas_rows[0])
                workbook_data = workbook_export.export_xlsx()
                with ZipFile(io.BytesIO(workbook_data)) as archive:
                    ns = {"s": workbook_export.NS}
                    main = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                    atlas = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
                    main_cells = {cell.get("r"): "".join(cell.itertext())
                                  for cell in main.findall("s:sheetData/s:row/s:c", ns)}
                    atlas_cells = {cell.get("r"): "".join(cell.itertext())
                                   for cell in atlas.findall("s:sheetData/s:row/s:c", ns)}
                    atlas_xml_cells = {cell.get("r"): cell
                                       for cell in atlas.findall("s:sheetData/s:row/s:c", ns)}
                    atlas_headers = list(atlas_rows[0])
                    last_column = workbook_export._column(len(atlas_headers))
                    self.assertEqual(atlas.find("s:dimension", ns).get("ref"), f"A1:{last_column}3")
                    self.assertEqual(len(atlas.findall("s:sheetData/s:row", ns)), 3)
                    number_column = workbook_export._column(atlas_headers.index(choice_column) + 1)
                    allocation_column = workbook_export._column(atlas_headers.index(choice_allocation_column) + 1)
                    for row_number, row in enumerate(atlas_rows, 2):
                        cell = atlas_xml_cells[f"{number_column}{row_number}"]
                        self.assertIsNone(cell.get("t"))
                        self.assertEqual(cell.findtext("s:v", namespaces=ns), row[choice_column])
                        self.assertEqual(atlas_cells[f"{allocation_column}{row_number}"],
                                         row[choice_allocation_column])
                    links = main.findall("s:hyperlinks/s:hyperlink", ns)
                    self.assertEqual(len(links), len(main_rows))
                    for link in links:
                        self.assertTrue(link.get("location").startswith("'Atlas Character Settings'!"))
                        target = link.get("location").rsplit("!", 1)[-1]
                        self.assertEqual(main_cells[link.get("ref")], atlas_cells[target])
            finally:
                store.DATA_DIR = previous
                logger._READY = False

    def test_latest_currency_counts_extend_recipe_rows_once_and_keep_scan_history(self):
        from PoE2_Data_Logger.core import atlas_catalog, catalog_repairs
        with tempfile.TemporaryDirectory(prefix="atlas-workbook-currency-") as directory:
            previous = store.DATA_DIR
            try:
                store.DATA_DIR = Path(directory)
                logger._READY = False
                logger.initialize()
                logger.clear_export_and_reset_ids()
                logger.start_map()
                logger.save_atlas_settings({"catalog_version": atlas_catalog.catalog()["version"],
                    "allocated": [], "choices": {}, "gear_item_rarity": 37.5})
                modifiers = ["25% increased Rarity of Items found", "Monsters deal 20% increased Damage"]
                logger.save_settings({"waystone_mods": modifiers})
                logger.commit_remnant(catalog_repairs.FARRUL_GRACE, catalog_repairs.FARRUL_HUNT, 47)
                logger.save_currency_snapshot("end", [{"name": "Chaos Orb", "quantity": 3}])
                logger.save_currency_snapshot("end", [{"name": "Chaos Orb", "quantity": 7},
                    {"name": "Reviewed fragment", "quantity": 2}], register_names=True)
                logger.finish_map(0, 0, 0, 0)
                logger.start_map()
                logger.save_currency_snapshot("end", [{"name": "Chaos Orb", "quantity": 4}])

                primary = csv.DictReader(io.StringIO(logger.export_primary_csv().decode("utf-8-sig")))
                headers = primary.fieldnames
                rows = list(primary)
                self.assertEqual(headers[:4], ["Recipe Input", "Matched Recipe", "Socket Count", "Exact Rune Combo"])
                self.assertTrue(all(header not in headers for header in
                    ("Map Mods", "# +2 Mod Tablets", "Tablet Mods", "Total Mods", "Master +Mods",
                     "Base Map Mods", "Waystone %", "Waystone Modifiers")))
                first_map = [row for row in rows if row["Map ID"] == "M0001"]
                self.assertEqual([row["Matched Recipe"] for row in first_map],
                                 [catalog_repairs.FARRUL_GRACE, catalog_repairs.FARRUL_HUNT])
                self.assertEqual([row["Socket Count"] for row in first_map], ["5", "5"])
                self.assertTrue(all(row["Exact Rune Combo"] for row in first_map))
                self.assertEqual([row["Chaos Orb"] for row in first_map], ["7", ""])
                self.assertEqual([row["Reviewed fragment"] for row in first_map], ["2", ""])
                self.assertEqual([row["Omen of Bartering"] for row in first_map], ["0", ""])
                self.assertEqual([row["Map Modifiers"] for row in first_map], ["\n".join(modifiers), ""])
                second_map = [row for row in rows if row["Map ID"] == "M0002"]
                self.assertEqual(len(second_map), 1)
                self.assertEqual(second_map[0]["Matched Recipe"], "")
                self.assertEqual(second_map[0]["Chaos Orb"], "4")
                self.assertEqual(second_map[0]["Start Baseline"], "Assumed empty")

                history_rows = list(csv.DictReader(io.StringIO(logger.export_all_csv().decode("utf-8-sig"))))
                chaos_history = [row for row in history_rows
                                 if row["Type"] == "Currency" and row["Currency"] == "Chaos Orb"]
                self.assertEqual([row["Quantity"] for row in chaos_history], ["3", "7", "4"])
                workbook_data = workbook_export.export_xlsx()
                with ZipFile(io.BytesIO(workbook_data)) as archive:
                    ns = {"s": workbook_export.NS}
                    main = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                    xml_rows = main.findall("s:sheetData/s:row", ns)
                    self.assertEqual(len(xml_rows), len(rows) + 1)
                    actual_headers = ["".join(cell.itertext()) for cell in xml_rows[0]]
                    self.assertEqual(actual_headers, headers)
                    cells = {cell.get("r"): cell for row in xml_rows[1:] for cell in row}
                    for field, first_count in (("Chaos Orb", "7"), ("Reviewed fragment", "2")):
                        column = workbook_export._column(headers.index(field) + 1)
                        cell = cells[f"{column}2"]
                        self.assertIsNone(cell.get("t"))
                        self.assertEqual(cell.findtext("s:v", namespaces=ns), first_count)
                        self.assertNotIn(f"{column}3", cells)
                    history = ET.fromstring(archive.read("xl/worksheets/sheet3.xml"))
                    self.assertEqual(len(history.findall("s:sheetData/s:row", ns)), len(history_rows) + 1)
                    links = main.findall("s:hyperlinks/s:hyperlink", ns)
                    self.assertEqual(len(links), len(rows))
                    self.assertTrue(all(link.get("location") == "'Atlas Character Settings'!A2" for link in links))
                if importlib.util.find_spec("openpyxl"):
                    from openpyxl import load_workbook
                    workbook = load_workbook(io.BytesIO(workbook_data), data_only=True)
                    try:
                        sheet = workbook["Export"]
                        reader_headers = [cell.value for cell in sheet[1]]
                        read_rows = [dict(zip(reader_headers, values))
                                     for values in sheet.iter_rows(min_row=2, values_only=True)]
                        read_first_map = [row for row in read_rows if row["Map ID"] == "M0001"]
                        self.assertEqual([row["Matched Recipe"] for row in read_first_map],
                                         [catalog_repairs.FARRUL_GRACE, catalog_repairs.FARRUL_HUNT])
                        self.assertEqual([row["Chaos Orb"] for row in read_first_map], [7, None])
                        self.assertEqual([row["Reviewed fragment"] for row in read_first_map], [2, None])
                        count_cell = sheet.cell(2, reader_headers.index("Chaos Orb") + 1)
                        self.assertEqual(count_cell.data_type, "n")
                    finally:
                        workbook.close()
            finally:
                store.DATA_DIR = previous
                logger._READY = False

    def test_invalid_atlas_row_width_rejected(self):
        with patch.object(logger, "_connect", memory_database), \
                patch.object(logger, "export_primary_csv", return_value=csv_bytes(["Map ID"], [])), \
                patch.object(logger, "export_all_csv", return_value=csv_bytes(["Map ID"], [])), \
                patch.object(logger, "export_atlas_csv",
                             return_value=csv_bytes(ATLAS_HEADERS, [["too", "short"]])), \
                patch.object(logger, "map_summary_item_headers", return_value=[]):
            with self.assertRaisesRegex(ValueError, "invalid width"):
                workbook_export.export_xlsx()

    def test_all_sheets_use_the_same_read_snapshot_during_a_concurrent_edit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "export.sqlite3"
            writer = sqlite3.connect(path)
            try:
                writer.execute("PRAGMA journal_mode=WAL")
                writer.execute("CREATE TABLE setup (rarity INTEGER)")
                writer.execute("INSERT INTO setup VALUES (25)")
                writer.commit()

                @contextmanager
                def reader_database():
                    reader = sqlite3.connect(path)
                    try:
                        yield reader
                    finally:
                        reader.close()

                def main_export(*, _db):
                    rarity = _db.execute("SELECT rarity FROM setup").fetchone()[0]
                    writer.execute("UPDATE setup SET rarity=75")
                    writer.commit()
                    return csv_bytes(["Gear Item Rarity %"], [[rarity]])

                def atlas_export(*, _db):
                    rarity = _db.execute("SELECT rarity FROM setup").fetchone()[0]
                    return csv_bytes(["Gear Item Rarity %"], [[rarity]])

                with patch.object(logger, "_connect", reader_database), \
                        patch.object(logger, "export_primary_csv", side_effect=main_export), \
                        patch.object(logger, "export_atlas_csv", side_effect=atlas_export), \
                        patch.object(logger, "export_all_csv", side_effect=atlas_export), \
                        patch.object(logger, "map_summary_item_headers", return_value=[]):
                    data = workbook_export.export_xlsx()
                self.assertEqual(writer.execute("SELECT rarity FROM setup").fetchone()[0], 75)
                with ZipFile(io.BytesIO(data)) as archive:
                    for number in (1, 2, 3):
                        sheet = ET.fromstring(archive.read(f"xl/worksheets/sheet{number}.xml"))
                        cell = sheet.find(f".//{{{workbook_export.NS}}}c[@r='A2']")
                        self.assertEqual(cell.findtext(f"{{{workbook_export.NS}}}v"), "25")
            finally:
                writer.close()

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional Excel reader is not installed")
    def test_excel_reader_opens_all_sheets_with_numeric_and_blank_values(self):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(self.sample_workbook()), data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Export", "Atlas Character Settings", "Scan History"])
            self.assertEqual(workbook["Export"]["C2"].value, 0)
            self.assertEqual(workbook["Export"]["B2"].hyperlink.location,
                             "'Atlas Character Settings'!A2")
            self.assertIsNone(workbook["Export"]["B4"].hyperlink)
            self.assertEqual(workbook["Export"]["D2"].value, 7)
            self.assertEqual(workbook["Export"]["D2"].data_type, "n")
            self.assertEqual(workbook["Export"]["D3"].value, 0)
            self.assertEqual(workbook["Scan History"]["D2"].value, 3)
            self.assertEqual(workbook["Scan History"]["D3"].value, 7)
            self.assertEqual(workbook["Scan History"]["B3"].hyperlink.location,
                             "'Atlas Character Settings'!A2")
            sheet = workbook["Atlas Character Settings"]
            self.assertEqual(sheet["D2"].value, 0)
            self.assertEqual(sheet["H2"].value, 1)
            self.assertEqual(sheet["H3"].value, 2)
            self.assertEqual(sheet["H2"].data_type, "n")
            self.assertEqual(sheet["H3"].data_type, "n")
            self.assertEqual(sheet["D3"].value, 175.5)
            self.assertIsNone(sheet["D4"].value)
            self.assertIsNone(sheet["G4"].value)
            self.assertIsNone(sheet["H4"].value)
            self.assertEqual(sheet["F2"].value, "No")
            self.assertEqual(sheet["G2"].value, "Yes")
            self.assertEqual(sheet["G3"].value, "No")
            self.assertEqual(sheet["E2"].value, "000123")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:H4")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()

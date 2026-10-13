"""Spreadsheet export auto-saves valid Atlas edits through the existing save action."""

import csv
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox, QPushButton

from PoE2_Data_Logger.core import logger_store as logger, store
from PoE2_Data_Logger.ui.native_desktop import LoggerWindow


class ExportSettingsGuardTests(unittest.TestCase):
    """Exercise export buttons and menu actions with valid and incomplete Atlas drafts."""
    @classmethod
    def setUpClass(cls):
        """Create or reuse the QApplication required by the Qt test fixtures."""
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        """Open an isolated logger with saved rarity, inline export tasks and a temporary export folder."""
        self.temporary = tempfile.TemporaryDirectory(prefix="poe2-export-settings-")
        self.previous_data = store.DATA_DIR
        store.DATA_DIR = Path(self.temporary.name)
        logger._READY = False
        logger.initialize()
        logger.start_map()
        logger.save_atlas_settings({"allocated": [], "choices": {}, "gear_item_rarity": 12})
        self.window = LoggerWindow()
        self.window._poll.stop()
        self.window.show()
        self.page = self.window.atlas_settings_page
        self.submissions = []
        self.results = []
        self.window._submit = self.submit
        self.folder = Path(self.temporary.name) / "exports"
        self.folder.mkdir()
        self.window.export_folder.setText(str(self.folder))

    def tearDown(self):
        """Close the logger window, stop workers and restore the original data directory."""
        self.window.close()
        self.window.pool.shutdown(wait=True, cancel_futures=True)
        self.app.processEvents()
        store.DATA_DIR = self.previous_data
        logger._READY = False
        self.temporary.cleanup()

    def submit(self, label, work, done):
        """Record and run an export task inline, retaining its result before invoking the callback."""
        self.submissions.append(label)
        result = work()
        self.results.append(result)
        done(result)

    def edit_rarity(self):
        """Type a new rarity draft and assert it differs from the persisted Atlas setting."""
        self.window.tabs.setCurrentWidget(self.page)
        self.app.processEvents()
        field = self.page.gear_rarity.lineEdit()
        field.selectAll()
        QTest.keyClicks(field, "187.25")
        self.assertTrue(self.page.dirty)
        self.assertEqual(self.page.settings()["gear_item_rarity"], 187.25)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def trigger_export(self, entrypoint, kind, *, navigate=True):
        """Activate the requested folder button or Save As menu action for an export format."""
        if navigate:
            self.window.tabs.setCurrentIndex(5)
        if entrypoint == "folder":
            self.window.export_folder.setText(str(self.folder))
            button = next(button for button in self.window.findChildren(QPushButton)
                          if button.text() == f"Save {kind.upper()} to folder")
            button.click()
        else:
            menu_action = self.window.menuBar().actions()[0]
            menu = menu_action.menu()
            action = next(action for action in menu.actions()
                          if action.text() == f"Save Export {kind.upper()} as…")
            action.trigger()

    def assert_exported_rarity(self, kind, expected):
        """Reopen the generated CSV or workbook and verify its saved gear rarity."""
        destination = Path(self.results[-1]["path"])
        self.assertTrue(destination.is_file())
        if kind == "csv":
            with destination.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertTrue(rows)
            self.assertEqual({row["Gear Item Rarity %"] for row in rows}, {str(expected)})
            atlas_path = Path(self.results[-1]["atlas_path"])
            history_path = Path(self.results[-1]["history_path"])
            self.assertTrue(atlas_path.is_file())
            self.assertTrue(history_path.is_file())
            with atlas_path.open(encoding="utf-8-sig", newline="") as stream:
                atlas_rows = list(csv.DictReader(stream))
            self.assertIn(str(expected), {row["Gear Item Rarity %"] for row in atlas_rows})
        else:
            with ZipFile(destination) as archive:
                for number in (1, 2):
                    sheet = ET.fromstring(archive.read(f"xl/worksheets/sheet{number}.xml"))
                    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    rows = sheet.findall("s:sheetData/s:row", ns)
                    headers = {"".join(cell.itertext()): re.sub(r"\d+$", "", cell.get("r"))
                               for cell in rows[0].findall("s:c", ns)}
                    column = headers["Gear Item Rarity %"]
                    values = {"".join(cell.itertext()) for row in rows[1:]
                              for cell in row.findall("s:c", ns)
                              if re.sub(r"\d+$", "", cell.get("r")) == column}
                    if number == 1:
                        self.assertEqual(values, {str(expected)})
                    else:
                        self.assertIn(str(expected), values)
            workbook = load_workbook(destination, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames,
                                 ["Export", "Atlas Character Settings", "Scan History", "Analysis Context",
                                  "Family Database", "Recipe Database", "Seed Database", "Reference Aliases"])
                for name in workbook.sheetnames:
                    rows = workbook[name].iter_rows(values_only=True)
                    headers = next(rows)
                    if name in ("Export", "Atlas Character Settings"):
                        column = headers.index("Gear Item Rarity %")
                        values = {row[column] for row in rows}
                        self.assertIn(float(expected), values)
                    else:
                        # The independent reader must parse every sheet's rows,
                        # including Context/reference data, rather than only its title.
                        list(rows)
            finally:
                workbook.close()

    def check_export_draft(self, entrypoint, kind):
        """Auto-save the current Atlas draft and export it through the real UI action."""
        self.edit_rarity()
        self.page.toggle_node("AtlasExpeditionNotable8")
        draft = self.page.settings()
        before = logger.get_state()
        destination = self.folder / f"chosen.{kind}"
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")) as chooser, \
                patch.object(QMessageBox, "question") as question:
            self.trigger_export(entrypoint, kind)
            question.assert_not_called()
            self.assertEqual(len(self.submissions), 1)
            self.assertFalse(self.page.dirty)
            self.assertEqual(self.window.tabs.currentIndex(), 5)
            self.assertIn("Export saved:", self.window.statusBar().currentMessage())
            self.assertEqual(self.page.settings(), draft)
            after = logger.get_state()
            self.assertEqual({key: after["settings"]["atlas_settings"][key] for key in draft}, draft)
            self.assertEqual(after["settings"]["atlas_settings"]["catalog_id"],
                             before["settings"]["atlas_settings"]["catalog_id"])
            self.assertEqual(after["scan_commit_count"], before["scan_commit_count"] + 1)
            self.assertEqual(after["current_map_id"], before["current_map_id"])
            if entrypoint == "save_as":
                chooser.assert_called_once()
            else:
                chooser.assert_not_called()
            question.assert_not_called()
        self.assert_exported_rarity(kind, 187.25)

    def test_folder_csv_auto_saves_atlas_without_redirecting(self):
        """Verify the CSV folder button saves the draft before producing all three CSVs."""
        self.check_export_draft("folder", "csv")

    def test_folder_xlsx_auto_saves_atlas_without_redirecting(self):
        """Verify the Excel folder button saves the draft before producing a workbook."""
        self.check_export_draft("folder", "xlsx")

    def test_save_as_csv_auto_saves_atlas_without_redirecting(self):
        """Verify the CSV menu saves the draft before producing its companion files."""
        self.check_export_draft("save_as", "csv")

    def test_save_as_xlsx_auto_saves_atlas_without_redirecting(self):
        """Verify the Excel menu saves the draft before producing a readable workbook."""
        self.check_export_draft("save_as", "xlsx")

    def check_incomplete_rarity(self, kind):
        """Reject an unfinished draft without silently exporting its earlier numeric value."""
        self.edit_rarity()
        field = self.page.gear_rarity.lineEdit()
        field.selectAll()
        QTest.keyClicks(field, ".")
        raw_text = field.text()
        before = logger.get_state()
        destination = self.folder / f"chosen.{kind}"
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
            self.trigger_export("save_as", kind, navigate=False)
        self.assertEqual(self.submissions, [])
        self.assertEqual(list(self.folder.iterdir()), [])
        self.assertEqual(self.window.tabs.currentWidget(), self.page)
        self.assertTrue(self.page.dirty)
        self.assertEqual(field.text(), raw_text)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"], before["settings"]["atlas_settings"])
        self.assertEqual(logger.get_state()["scan_commit_count"], before["scan_commit_count"])
        self.assertIn("gear item rarity", self.page.status.text().lower())

    def test_save_as_csv_rejects_incomplete_rarity(self):
        """Prevent CSV export from saving or discarding unfinished Atlas numeric text."""
        self.check_incomplete_rarity("csv")

    def test_save_as_xlsx_rejects_incomplete_rarity(self):
        """Prevent Excel export from saving or discarding unfinished Atlas numeric text."""
        self.check_incomplete_rarity("xlsx")

    def test_cancelled_xlsx_keeps_page_and_unsaved_atlas(self):
        """Cancel the Excel chooser without writing files, redirecting or changing the draft."""
        self.edit_rarity()
        draft = self.page.settings()
        before = logger.get_state()["scan_commit_count"]
        with patch.object(QFileDialog, "getSaveFileName", return_value=("", "")) as chooser:
            self.trigger_export("save_as", "xlsx")
        chooser.assert_called_once()
        self.assertEqual(self.window.tabs.currentIndex(), 5)
        self.assertEqual(self.submissions, [])
        self.assertEqual(list(self.folder.iterdir()), [])
        self.assertTrue(self.page.dirty)
        self.assertEqual(self.page.settings(), draft)
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_cancelled_csv_does_not_save_atlas(self):
        """Cancel the CSV chooser without adding an Atlas commit or consuming its draft."""
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        with patch.object(QFileDialog, "getSaveFileName", return_value=("", "")):
            self.trigger_export("save_as", "csv")
        self.assertEqual(self.submissions, [])
        self.assertTrue(self.page.dirty)
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_declined_csv_companion_does_not_save_atlas(self):
        """Keep edits pending when the user declines replacing an existing CSV companion."""
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        destination = self.folder / "chosen.csv"
        companion = self.folder / "chosen_Atlas.csv"
        companion.write_bytes(b"existing Atlas data")
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")), \
                patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            self.trigger_export("save_as", "csv")
        self.assertEqual(companion.read_bytes(), b"existing Atlas data")
        self.assertFalse(destination.exists())
        self.assertEqual(self.submissions, [])
        self.assertTrue(self.page.dirty)
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_invalid_folder_does_not_auto_save_atlas(self):
        """Reject a missing export directory before saving either spreadsheet's Atlas draft."""
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        self.window.tabs.setCurrentIndex(5)
        self.window.export_folder.setText(str(self.folder / "missing"))
        for kind in ("csv", "xlsx"):
            button = next(button for button in self.window.findChildren(QPushButton)
                          if button.text() == f"Save {kind.upper()} to folder")
            button.click()
            self.assertEqual(self.submissions, [])
            self.assertEqual(list(self.folder.iterdir()), [])
            self.assertTrue(self.page.dirty)
            self.assertEqual(logger.get_state()["scan_commit_count"], before)
            self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_clean_and_repeated_exports_do_not_duplicate_atlas_commits(self):
        """Do not create Atlas saves for a clean editor or for a draft already auto-saved."""
        before = logger.get_state()["scan_commit_count"]
        for kind in ("csv", "xlsx"):
            self.trigger_export("folder", kind)
            self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.edit_rarity()
        self.trigger_export("folder", "xlsx")
        self.assertEqual(logger.get_state()["scan_commit_count"], before + 1)
        self.trigger_export("folder", "csv")
        self.assertEqual(logger.get_state()["scan_commit_count"], before + 1)

    def test_explicit_atlas_save_does_not_repeat_when_exporting(self):
        """Respect a draft already saved through the existing Atlas save button."""
        self.edit_rarity()
        self.page.save_button.click()
        before = logger.get_state()["scan_commit_count"]
        self.trigger_export("folder", "xlsx")
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assert_exported_rarity("xlsx", 187.25)

    def check_real_worker_export(self, kind):
        """Auto-save a draft then complete the real export worker and Qt callback."""
        self.edit_rarity()
        draft = self.page.settings()
        before = logger.get_state()["scan_commit_count"]
        real_submit = LoggerWindow._submit.__get__(self.window, LoggerWindow)
        real_done = self.window._export_saved

        def record_done(result):
            """Retain the actual worker result and deliver the production export completion."""
            self.results.append(result)
            real_done(result)

        with patch.object(self.window, "_submit", real_submit), \
                patch.object(self.window, "_export_saved", side_effect=record_done):
            self.trigger_export("folder", kind)
            deadline = time.monotonic() + 15
            while self.window._pending_tasks and time.monotonic() < deadline:
                self.app.processEvents()
                time.sleep(0.01)
        self.assertEqual(self.window._pending_tasks, {}, self.window.statusBar().currentMessage())
        self.assertIn("Export saved:", self.window.statusBar().currentMessage())
        self.assertEqual(self.window.tabs.currentIndex(), 5)
        self.assertFalse(self.page.dirty)
        self.assertEqual(self.page.settings(), draft)
        saved = logger.get_state()["settings"]["atlas_settings"]
        self.assertEqual({key: saved[key] for key in draft}, draft)
        self.assertEqual(logger.get_state()["scan_commit_count"], before + 1)
        self.assertEqual(len(list(self.folder.iterdir())), 3 if kind == "csv" else 1)
        self.assertEqual(len(self.results), 1)
        self.assert_exported_rarity(kind, 187.25)

    def test_folder_xlsx_real_worker_creates_readable_workbook_without_redirecting(self):
        """Create and independently reopen all Excel sheets through the real worker and Qt callback."""
        self.check_real_worker_export("xlsx")

    def test_folder_csv_real_worker_creates_companions_without_redirecting(self):
        """Create all three CSV files through the real worker and Qt callback with an Atlas draft."""
        self.check_real_worker_export("csv")

    def test_database_backup_remains_available_with_unsaved_settings(self):
        """The SQL Save As path neither auto-saves Atlas nor changes its backup behavior."""
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        destination = self.folder / "backup.sqlite3"
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
            self.window.save_as("backup")
        self.assertEqual(len(self.submissions), 1)
        self.assertTrue(destination.is_file())
        self.assertTrue(self.page.dirty)
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_database_folder_backup_does_not_auto_save_atlas(self):
        """The raw SQL folder button keeps pending Atlas edits outside the backup."""
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        self.window.tabs.setCurrentIndex(5)
        self.window.raw_database_export_button.click()
        self.assertEqual(len(self.submissions), 1)
        self.assertTrue(Path(self.results[-1]["path"]).is_file())
        self.assertTrue(self.page.dirty)
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)

    def test_auto_save_does_not_rewrite_frozen_map_activity(self):
        """Reuse normal Atlas saving: keep current-map activity frozen and stage the edited setup."""
        logger.save_counts(10, 2, 1, unique=0)
        self.window.refresh()
        with logger._connect() as db:
            frozen = db.execute("SELECT snapshot_json FROM maps WHERE map_id='M0001'").fetchone()[0]
        self.edit_rarity()
        before = logger.get_state()["scan_commit_count"]
        self.trigger_export("folder", "xlsx")
        after = logger.get_state()
        self.assertFalse(self.page.dirty)
        self.assertEqual(after["settings"]["atlas_settings"]["gear_item_rarity"], 187.25)
        self.assertEqual(after["atlas_settings_target_map_id"], "M0002")
        self.assertEqual(after["current_map_id"], "M0001")
        self.assertEqual(after["scan_commit_count"], before + 1)
        with logger._connect() as db:
            self.assertEqual(db.execute("SELECT snapshot_json FROM maps WHERE map_id='M0001'").fetchone()[0], frozen)
        workbook = load_workbook(self.results[-1]["path"], read_only=True)
        try:
            rows = workbook["Export"].iter_rows(values_only=True)
            headers = next(rows)
            map_column, rarity_column = headers.index("Map ID"), headers.index("Gear Item Rarity %")
            rarities = {}
            for row in rows:
                rarities.setdefault(row[map_column], set()).add(row[rarity_column])
            self.assertEqual(rarities["M0001"], {12})
            self.assertEqual(rarities["M0002"], {187.25})
        finally:
            workbook.close()

    def test_real_writer_save_failure_retains_draft_and_allows_sql_save_as(self):
        """A real SQLite writer blocks auto-saving quickly without producing a stale spreadsheet."""
        self.edit_rarity()
        draft = self.page.settings()
        before = logger.get_state()["scan_commit_count"]
        writer = sqlite3.connect(store.DATA_DIR / "scans.sqlite3", timeout=0)
        try:
            writer.execute("BEGIN IMMEDIATE")
            for kind in ("csv", "xlsx"):
                destination = self.folder / f"blocked.{kind}"
                started = time.monotonic()
                with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
                    self.trigger_export("save_as", kind)
                self.assertLess(time.monotonic() - started, 1)
                self.assertEqual(self.submissions, [])
                self.assertFalse(destination.exists())
                self.assertTrue(self.page.dirty)
                self.assertEqual(self.page.settings(), draft)
                self.assertEqual(self.window.tabs.currentIndex(), 5)
                self.assertEqual(logger.get_state()["scan_commit_count"], before)
            backup = self.folder / "during-writer.sqlite3"
            with patch.object(QFileDialog, "getSaveFileName", return_value=(str(backup), "")):
                self.window.run(lambda: self.window.save_as("backup"))
            self.assertTrue(backup.is_file())
            self.assertTrue(self.page.dirty)
        finally:
            writer.rollback()
            writer.close()
        self.assertEqual(logger.get_state()["scan_commit_count"], before)
        self.assertEqual(logger.get_state()["settings"]["atlas_settings"]["gear_item_rarity"], 12)


if __name__ == "__main__":
    unittest.main()

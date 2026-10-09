"""Opt-in end-user release gate, using Qt controls and an independent ledger.

Run: QT_QPA_PLATFORM=offscreen POE2_RUN_HUNDRED_MAP_SESSION=1 \
     .venv/bin/python -m unittest tests.test_hundred_map_session -v

POE2_SESSION_ARTIFACTS selects a new output directory (default: /tmp/poe2-100-map-*).
Synthetic OCR results enter the real capture completion/read APIs. OCR accuracy,
Windows hotkeys, physical screen capture and game integration are separate gates.
No user action calls logger.save_*, logger.commit_* or SQL writes directly.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import copy
import csv
import faulthandler
import hashlib
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np
from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QPushButton

from PoE2_Data_Logger.core import logger_store as logger, store
from PoE2_Data_Logger.ocr import item_ocr
from PoE2_Data_Logger.ui.native_desktop import LoggerWindow


OPENED = [
    {"recipe": "Divine Orb x3", "sockets": 6,
     "combo": "Death + Soul + Power + Life + Vision + Celestial"},
    {"recipe": "Divine Orb x2", "sockets": 5,
     "combo": "Death + Soul + Power + Life + Vision"},
    {"recipe": "Divine Orb", "sockets": 4,
     "combo": "Death + Soul + Power + Life"},
]
PERKS = {"Jado": ["Unexpected Missions", "Long Days", "None", "None"],
         "Doryani": ["Careful Procurement", "Hidden Patterns", "None", "None"],
         "Hilda": ["Ancient Inscriptions", "Dangerous Game", "None", "None"]}
WAYSTONE_FIELDS = sorted(["waystone", "tier", "map_mods", "item_rarity", "monster_rarity",
                         "pack_size", "effectiveness", "waystone_name", "waystone_mods"])


@unittest.skipUnless(os.getenv("POE2_RUN_HUNDRED_MAP_SESSION") == "1", "opt-in 100-map GUI gate")
class HundredMapSession(unittest.TestCase):
    """Exercise an opt-in 100-map Qt workflow against an independent expected ledger."""
    @classmethod
    def setUpClass(cls):
        """Create the shared Qt application for the end-user session gate."""
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        """Open an isolated logger and initialize the expected ledger and captured artwork."""
        directory = os.getenv("POE2_SESSION_ARTIFACTS")
        self.output = Path(directory) if directory else Path(tempfile.mkdtemp(prefix="poe2-100-map-"))
        self.output.mkdir(parents=True, exist_ok=True)
        self.previous = store.DATA_DIR
        store.DATA_DIR = self.output / "data"
        logger._READY = False
        logger.initialize()
        self.w = LoggerWindow()
        self.w._poll.stop()
        self.w.resize(1800, 1200)
        self.w.show()
        self.app.processEvents()
        self.started = time.monotonic()
        self.failures = []
        self.checks = Counter()
        self.coverage = Counter()
        self.expected = {}
        self.commit_log = []
        self.frozen = {}
        self.labels = set()
        self.rejected_labels = set()
        self.exports = []
        self.completed = 0
        self.current = 1
        self.expedition = 1
        self.selections = {master: ["None"] * 4 for master in PERKS}
        self.atlas = {"catalog_version": self.w.atlas_settings_page._catalog["version"],
                      "catalog_id": hashlib.sha256(json.dumps(self.w.atlas_settings_page._catalog,
                          ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest(),
                      "allocated": [], "choices": {}, "gear_item_rarity": None}
        self.artwork = Image.new("RGB", (648, 270), "black")
        tile = Image.fromarray(np.random.default_rng(907).integers(20, 235, (54, 54, 3), dtype=np.uint8))
        self.artwork.paste(tile, (0, 0))
        self.artwork.info["poe2_inventory_aligned"] = True
        self.raw = io.BytesIO()
        self.artwork.save(self.raw, format="PNG")

    def tearDown(self):
        """Write session evidence and metrics, close workers and restore the data directory."""
        metrics = {"completed_maps": self.completed, "checks": dict(self.checks),
                   "coverage": dict(self.coverage), "expected_commits": len(self.commit_log),
                   "failures": self.failures, "duration_seconds": round(time.monotonic() - self.started, 2),
                   "artifacts": str(self.output), "exports": self.exports,
                   "limits": ["OCR outputs injected at scan ingress; OCR accuracy is a separate gate",
                              "Qt offscreen: no Windows game capture/hotkey or physical pointer validation",
                              "Database backup export and restart covered; no database backup import UI exists"]}
        (self.output / "metrics.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
        (self.output / "expected_ledger.json").write_text(json.dumps(self.expected, indent=2), encoding="utf-8")
        print("100-map metrics: " + str(self.output / "metrics.json"), flush=True)
        self.w.close()
        self.w.pool.shutdown(wait=True, cancel_futures=True)
        self.app.processEvents()
        store.DATA_DIR = self.previous
        logger._READY = False

    def check(self, actual, expected, context):
        """Count a contextual comparison and record mismatches without stopping the session."""
        self.checks[context.split(":", 1)[0]] += 1
        if actual != expected:
            self.failures.append({"map": self.current, "context": context,
                                  "actual": actual, "expected": expected})

    def wait_tasks(self):
        """Pump Qt until background work finishes or write thread diagnostics on timeout."""
        deadline = time.monotonic() + 90
        while self.w._pending_tasks and time.monotonic() < deadline:
            self.app.processEvents()
            # qWait's C++ event pump can retain the Python GIL while waiting,
            # throttling the exporter and its Python completion callback.
            # Keep Qt responsive, then explicitly yield to the worker.
            time.sleep(.005)
        if self.w._pending_tasks:
            details = {"map": self.current, "context": "background-task-timeout",
                       "pending_tasks": list(self.w._pending_tasks),
                       "status": self.w.statusBar().currentMessage()}
            self.failures.append(details)
            diagnostic = self.output / "ui-task-timeout.txt"
            with diagnostic.open("w", encoding="utf-8") as stream:
                stream.write(json.dumps(details, indent=2) + "\n\n")
                stream.flush()
                faulthandler.dump_traceback(file=stream, all_threads=True)
            stack_details = diagnostic.read_text(encoding="utf-8", errors="replace")[:8000]
            raise RuntimeError(f"UI background task timed out; thread stacks: {diagnostic}\n{stack_details}")
        self.app.processEvents()

    def choose(self, control, wanted):
        """Select a control option by data or text and reject unavailable values."""
        index = control.findData(wanted)
        if index < 0:
            index = control.findText(str(wanted))
        if index < 0:
            raise RuntimeError(f"Control does not offer {wanted!r}")
        control.setCurrentIndex(index)

    def button(self, text, page=None):
        """Click an enabled matching UI button on the requested tab."""
        if page is not None:
            self.w.tabs.setCurrentIndex(page)
        matches = [b for b in self.w.findChildren(QPushButton) if b.text() == text and b.isEnabled()]
        visible = [b for b in matches if b.isVisible()]
        target = (visible or matches)[0] if matches else None
        if target is None:
            raise RuntimeError(f"No enabled UI button: {text}")
        target.click()
        self.app.processEvents()

    def record(self, kind, details=None, snapshot=None, reference=None, map_id=None, expedition=None,
               batch_remaining=0):
        """Compare the next database commit with ledger expectations and record its coverage."""
        number = len(self.commit_log) + 1
        with logger._connect() as db:
            row = db.execute("SELECT * FROM commits WHERE number=?", (number,)).fetchone()
            count = db.execute("SELECT count(*) FROM commits").fetchone()[0]
        # A confident propagation scan atomically saves its accepted count and
        # chain part. Verify both audit entries in order against the ledger.
        expected_count = number + batch_remaining
        self.check(count, expected_count, "commits:count")
        if not row or count != expected_count:
            raise RuntimeError(f"UI did not create expected {kind} commit: {self.w.statusBar().currentMessage()}")
        saved = dict(row)
        mid = map_id or f"M{self.current:04}"
        eid = expedition if expedition is not None else self.expedition
        for key, expected in (("number", number), ("kind", kind), ("map_id", mid),
                              ("expedition_id", f"{mid}-E{eid:02}" if kind not in ("Currency", "Ritual") else "")):
            self.check(saved[key], expected, f"commits:{kind}:{key}")
        if reference is not None:
            self.check(saved["reference"], reference, f"commits:{kind}:reference")
        actual_details = json.loads(saved["details_json"])
        for key, val in (details or {}).items():
            self.check(actual_details.get(key), val, f"commits:{kind}:details:{key}")
        if snapshot:
            actual_context = json.loads(saved["snapshot_json"])
            for key, val in snapshot.items():
                self.check(actual_context.get(key), val, f"snapshot:{kind}:{key}")
        self.commit_log.append({"number": number, "kind": kind, "map_id": mid,
                                "details": copy.deepcopy(details or {})})
        self.check(logger.get_state()["scan_commit_count"], expected_count, "commits:header-count")
        self.coverage[kind] += 1

    def atlas_edit(self, n):
        """Click an Atlas choice node, type gear rarity and save the expected setup."""
        page = self.w.atlas_settings_page
        self.w.tabs.setCurrentIndex(12)
        page.clear_button.click()
        chosen = "AtlasExpeditionNotable8"
        self.choose(page.activity_filter, "Expedition")
        self.app.processEvents()
        item = page.node_items[chosen]
        page.view.centerOn(item)
        self.app.processEvents()
        position = page.view.mapFromScene(item.scenePos())
        QTest.mouseClick(page.view.viewport(), Qt.MouseButton.LeftButton, pos=position)
        self.app.processEvents()
        self.choose(page.choice_combo, page._nodes[chosen]["choices"][((n - 1) // 5) % len(page._nodes[chosen]["choices"])] ["id"])
        gear = round(((n - 1) // 5) * 7.25, 2)
        if gear.is_integer():
            gear = int(gear)
        editor = page.gear_rarity.lineEdit()
        editor.setFocus()
        editor.selectAll()
        QTest.keyClicks(editor, str(gear))
        QTest.keyClick(editor, Qt.Key.Key_Tab)
        self.app.processEvents()
        self.atlas = {**self.atlas, "allocated": [chosen],
                      "choices": {chosen: page._nodes[chosen]["choices"][((n - 1) // 5) % len(page._nodes[chosen]["choices"])] ["id"]},
                      "gear_item_rarity": gear}
        self.check(page.settings(), {k: v for k, v in self.atlas.items() if k != "catalog_id"}, "atlas:typed-draft")
        page.save_button.click()
        self.record("Atlas settings")
        self.check(page.dirty, False, "atlas:saved-draft-clean")
        self.coverage["typed gear and graphics node click"] += 1

    def configure(self, n):
        # Atlas edits precede activity so they apply to the intended current map.
        # Save through the character UI on every map; repeating a setup also
        # checks that identical setups remain shared in the compact workbook.
        """Set map, master, tablet and waystone options through UI controls and scan ingress."""
        self.atlas_edit(n)
        self.w.tabs.setCurrentIndex(3)
        master = ["Jado", "Doryani", "Hilda", "None"][(n - 1) % 4]
        self.choose(self.w.master, master)
        self.button("Save active master", 3)
        self.record("Atlas Master")
        configured = master if master != "None" else "Jado"
        self.choose(self.w.perk_master, configured)
        for control, val in zip(self.w.perk_boxes, PERKS[configured]):
            self.choose(control, val)
        self.selections[configured] = PERKS[configured][:]
        self.button("Save perks", 3)
        self.record("Master perks", reference=configured)
        capacity = 1 + (n - 1) % 4
        self.w.tabs.setCurrentIndex(2)
        self.button("Clear tablet config", 2)
        random_affix = "Map has additional random Modifiers"
        if self.w.tablet_affixes[0].findData(random_affix) < 0:
            self.w.new_affix.setText(random_affix)
            self.button("Add to Affix DB", 2)
            self.coverage["UI affix registration"] += 1
        self.choose(self.w.tablets_used, capacity)
        tablet_pairs, tablet_values, random_counts = [], [], []
        for i, (a, v) in enumerate(zip(self.w.tablet_affixes, self.w.tablet_values)):
            active = i // 4 < capacity
            affix = "Pack Size" if active and i % 4 == 0 else random_affix if active and i % 4 == 1 else ""
            amount = n % 29 + i // 4 if affix == "Pack Size" else (n + i // 4) % 4 if affix else None
            self.choose(a, affix)
            v.setText(str(amount) if amount is not None else "")
            pair = {"affix": affix, "value": amount, "unit": "count" if affix == random_affix else "%"}
            tablet_pairs.append(pair)
            tablet_values += [affix, amount if affix and affix != random_affix else ""]
        for i in range(4):
            random_counts.append((n + i) % 4 if i < capacity else 0)
        self.button("Save tablet config", 2)
        self.record("Tablet config")
        raw_tablets = [[], [], [], []]
        if n % 10 == 0:
            before = len(self.commit_log)
            self.w._hover_item_read({"kind": "tablet", "source": "screen OCR",
                "ocr_rows": [{"text": "Unclear tablet inscription", "score": .2}],
                "mods": [], "matches": [], "uncertain": [], **logger.scan_context()}, self.raw.getvalue())
            self.check(logger.get_state()["scan_commit_count"], before, "auto:unclear-tablet-held")
            self.check(self.w.pending_review_kind, "tablet", "auto:unclear-tablet-review")
            self.w.reject_scan_button.click()
            self.coverage["unclear tablet held with auto-all enabled"] += 1
        for slot in range(1, capacity + 1):
            raw_tablets[slot - 1] = [f"{n % 29 + slot - 1}% increased Pack Size",
                f"Map has {(n + slot - 1) % 4} additional random Modifiers"]
            self.w._hover_item_read({"kind": "tablet", "source": "screen OCR",
                "ocr_rows": [{"text": text, "score": .99} for text in raw_tablets[slot - 1]],
                "mods": raw_tablets[slot - 1], "matches": [], "uncertain": [],
                **logger.scan_context()}, self.raw.getvalue())
            self.record("Tablet config", reference=f"Tablet {slot}")
            self.check(self.w.pending_review_kind, None, "auto:tablet-saved-review-cleared")
            self.check(logger.tablet_next_slot(), slot + 1, "auto:tablet-scan-order")
            self.coverage["clear live tablet automatic commit"] += 1
        config = {"tier": 15 + n % 2, "waystone": (n * 13) % 251,
                  "base_map_mods": n % 8, "irradiated": "Yes" if n % 3 == 0 else "No",
                  "ocean": "Yes" if n % 5 == 0 else "No", "deli": "Yes" if n % 2 == 0 else "No",
                  "wisp": "Yes" if n % 7 == 0 else "No", "aldur": ["None", "All +5", "Lucky"][n % 3],
                  "biome": ["Water", "Mountain", "Grass", "Forest", "Swamp", "Desert", "Ocean", "Island"][n % 8],
                  "city_type": ["None", "Faridun", "Ezomyte", "Vaal"][n % 4],
                  "item_rarity": n * 1.25, "monster_rarity": n % 60,
                  "pack_size": n % 30, "effectiveness": n % 20,
                  "waystone_name": "", "waystone_mods": [f"Session mod {n} A", f"Session mod {n} B"] + [""] * 8,
                  "master": master, "perks": [v if v != "None" else "" for v in self.selections.get(master, ["None"] * 4)],
                  "master_selections": copy.deepcopy(self.selections), "tablets_used": capacity,
                  "tablet_capacity": capacity, "tablet_affixes": tablet_pairs,
                  "tablet_values": tablet_values, "tablet_random_mods": random_counts,
                  "tablet_raw_mods": raw_tablets, "expedition": 1}
        for field, val in ((self.w.waystone, config["waystone"]), (self.w.map_mods, config["base_map_mods"]),
                           (self.w.item_rarity, config["item_rarity"]), (self.w.monster_rarity, config["monster_rarity"]),
                           (self.w.pack_size, config["pack_size"]), (self.w.effectiveness, config["effectiveness"])):
            field.setText(str(val))
        self.choose(self.w.tier, config["tier"])
        for field, key in ((self.w.irradiated, "irradiated"), (self.w.ocean, "ocean"),
                           (self.w.deli, "deli"), (self.w.wisp, "wisp")):
            changed = field.isChecked() != (config[key] == "Yes")
            field.setChecked(config[key] == "Yes")
            if changed:
                self.record("Map settings")
        for field, key in ((self.w.aldur, "aldur"), (self.w.biome, "biome"), (self.w.city_type, "city_type")):
            changed = field.currentData() != config[key]
            self.choose(field, config[key])
            if changed and key != "aldur":
                self.record("Map settings")
        for field, val in zip(self.w.waystone_mod_fields, config["waystone_mods"]):
            field.setText(val)
        config["master_adds_mod"] = int(master == "Jado")
        config["map_mods"] = config["base_map_mods"] + config["master_adds_mod"]
        config["area"] = 79 + (config["tier"] == 16) + (config["ocean"] == "Yes") + (config["irradiated"] == "Yes")
        config["tablets"] = sum(v == 2 for v in random_counts)
        config["tablet_mods"] = sum(random_counts)
        config["total_mods"] = config["map_mods"] + config["tablet_mods"]
        atlas_json = json.dumps(self.atlas, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        config.update(atlas_settings=copy.deepcopy(self.atlas), atlas_catalog_id=self.atlas["catalog_id"],
                      atlas_setup_id="A" + hashlib.sha256(atlas_json.encode()).hexdigest(),
                      gear_item_rarity=self.atlas["gear_item_rarity"])
        self.button("Save map settings", 2)
        self.record("Map settings", snapshot=config)
        self.expected[f"M{n:04}"] = {"snapshot": config, "inventories": {}, "ritual": [], "chain": [],
                                      "recipes": [], "kills": [100 + n, n % 31, n % 17, n % 5]}
        fields = {key: config[key] for key in
                  ("tier", "waystone", "item_rarity", "monster_rarity", "pack_size", "effectiveness")}
        fields["map_mods"] = 2
        mods = config["waystone_mods"][:2]
        result = {"kind": "waystone", "source": "screen OCR", "name": f"Session Waystone {n}",
                  "fields": fields, "mods": mods,
                  "ocr_rows": [{"text": text, "score": .99} for text in mods], **logger.scan_context()}
        if n % 10 == 0:
            before = len(self.commit_log)
            self.w._hover_item_read({**result, "fields": {**fields, "tier": None}}, self.raw.getvalue())
            self.check(logger.get_state()["scan_commit_count"], before, "auto:unclear-waystone-held")
            self.check(self.w.pending_review_kind, "waystone", "auto:unclear-waystone-review")
            self.w.reject_scan_button.click()
            self.coverage["unclear waystone held with auto-all enabled"] += 1
        self.w._hover_item_read(result, self.raw.getvalue())
        config.update(base_map_mods=2, map_mods=2 + config["master_adds_mod"],
                      total_mods=2 + config["master_adds_mod"] + config["tablet_mods"],
                      waystone_name=result["name"])
        self.record("Map settings", snapshot=config)
        self.check(self.w.pending_review_kind, None, "auto:waystone-saved-review-cleared")
        self.coverage["clear live waystone automatic commit"] += 1
        self.coverage["distinct map setup"] += 1
        return config

    def inventory(self, phase, amounts, *, custom=None, rejected=None, unnamed=False, automatic=False, live=False):
        """Capture controlled inventory results, review rows and verify the committed snapshot."""
        self.choose(self.w.inventory_phase, phase)
        before = len(self.commit_log)
        if automatic:
            # A grid contains positive stacks; zero amounts represent absent
            # icons and must not be manufactured into a supposedly clear read.
            amounts = {name: quantity for name, quantity in amounts.items() if quantity > 0}
        items = [{"slot": i + 2, "name": name, "quantity": qty}
                 for i, (name, qty) in enumerate(amounts.items())]
        unknown = []
        for slot, name in enumerate([custom, rejected, "" if unnamed else None], 1):
            if name is not None:
                unknown.append({"slot": slot, "candidate": "", "quantity": 1})
        result = {"items": items, "unknown": unknown}
        # Real capture bookkeeping and queued completion; only scanner output is substituted.
        with patch.object(item_ocr, "scan_inventory_grid", return_value=result):
            self.w._inventory_captured(self.artwork.copy(), live=automatic or live)
            self.wait_tasks()
        offset = len(items)
        for i, name in enumerate([val for val in [custom, rejected, "" if unnamed else None] if val is not None]):
            row = offset + i
            self.w.inventory_table.item(row, 1).setText(name)
            self.w.inventory_table.item(row, 2).setText("2" if name == custom else "999")
            controls = self.w.inventory_table.cellWidget(row, 3)
            if name == custom:
                controls.findChild(QPushButton, "approveCurrency").click()
                self.check(name in logger.inventory_names(), False if name not in self.labels else True,
                           "learning:row-button-does-not-register")
                self.coverage["individual row approve"] += 1
            else:
                controls.findChild(QPushButton, "rejectCurrency").click()
                self.rejected_labels.add(name) if name else None
                self.coverage["individual row reject"] += 1
        if not automatic:
            self.check(logger.get_state()["scan_commit_count"], before, "learning:row-buttons-no-commit")
            self.check(self.w.pending_review_kind, "currency", "learning:whole-scan-decision-pending")
        saved = dict(amounts)
        if custom:
            saved[custom] = 2
            self.labels.add(custom)
        if not automatic:
            self.w.approve_scan_button.click()
        context = self.expected[f"M{self.current:04}"]["snapshot"]
        context = {**context, "expedition": self.expedition}
        baseline = "Scanned" if phase == "start" or "start" in self.expected[f"M{self.current:04}"]["inventories"] else "Assumed empty"
        self.record("Currency", {"phase": phase, "items": saved, "start_baseline": baseline}, context,
                    reference=phase.title() + " inventory")
        self.expected[f"M{self.current:04}"]["inventories"][phase] = saved
        self.check(self.w.approve_scan_button.isEnabled(), False, "review:saved-inventory-lock")
        self.check(self.w.inventory_table.cellWidget(0, 3).isEnabled() if self.w.inventory_table.rowCount() else False,
                   False, "review:saved-row-lock")
        self.coverage["clear live currency automatic commit" if automatic else "currency manual review commit"] += 1
        if live and not automatic:
            self.coverage["unclear currency held with auto-all enabled"] += 1
        self.verify_totals()

    def reject_inventory(self):
        """Reject a synthetic inventory scan and verify it leaves counts and totals unchanged."""
        self.choose(self.w.inventory_phase, "end")
        self.w._inventory_read({"items": [{"slot": 1, "name": "Chaos Orb", "quantity": 99999}], "unknown": []}, live=False)
        before = len(self.commit_log)
        self.w.reject_scan_button.click()
        self.check(logger.get_state()["scan_commit_count"], before, "review:rejected-scan-no-commit")
        self.coverage["whole scan reject"] += 1
        self.verify_totals()

    def verify_totals(self):
        """Compare persisted snapshots, backend gains and visible session totals with the ledger."""
        total = Counter()
        counted = assumed = pending = 0
        for mid, entry in self.expected.items():
            phases = entry["inventories"]
            if "end" in phases:
                counted += 1
                assumed += "start" not in phases
                total.update({name: max(0, qty - phases.get("start", {}).get(name, 0))
                              for name, qty in phases["end"].items()})
            elif "start" in phases:
                pending += 1
            with logger._connect() as db:
                persisted = {row["phase"]: json.loads(row["items_json"]) for row in
                             db.execute("SELECT phase,items_json FROM currency_snapshots WHERE map_id=?", (mid,))}
            self.check(persisted, phases, "inventory:persisted-exact")
        expected = {k: v for k, v in total.items() if v > 0}
        totals = logger.session_currency_totals()
        self.check({r["name"]: r["quantity"] for r in totals["items"]}, expected, "counter:backend-rolling")
        self.check({name: card.quantity for name, card in self.w.session_currency.cards.items() if card.quantity > 0},
                   expected, "counter:visible-rolling")
        for key, val in (("maps_counted", counted), ("maps_assumed_empty", assumed), ("maps_pending_end", pending)):
            self.check(totals[key], val, "counter:" + key)

    def ritual(self, n):
        """Exercise uncertain reviewed and clear automatic Ritual pages with editable reward fields."""
        rows = [{"category": "Omen", "name": "Omen of Whittling", "quantity": 1 + n % 3,
                 "tribute": 300 + n, "source": f"Ritual map {n}", "deferred": n % 2 == 0},
                {"category": "Item", "name": f"Session reward {n % 7}", "quantity": 2 + n % 4,
                 "tribute": 900 + n, "source": f"Ritual map {n}", "deferred": False}]
        raw = f"Synthetic Ritual evidence {n}"
        result = {"items": rows + [{"name": "", "category": "Item", "quantity": 1, "tribute": 0, "deferred": False}],
                  "raw_text": raw, "tribute_available": 10000 + n, "rerolls_remaining": n % 5}
        before = len(self.commit_log)
        self.w._ritual_read(result, live=True)
        self.check(logger.get_state()["scan_commit_count"], before, "auto:unclear-ritual-held")
        self.check(self.w.pending_review_kind, "ritual", "auto:unclear-ritual-review")
        self.coverage["unclear Ritual held with auto-all enabled"] += 1
        # Editing actual counters and reward cells exercises the review form.
        self.w.ritual_tribute.setText(str(20000 + n))
        self.w.ritual_rerolls.setText(str(n % 4))
        self.w.ritual_table.item(1, 2).setText(str(rows[1]["quantity"] + 1))
        rows[1]["quantity"] += 1
        self.w.ritual_table.item(1, 5).setCheckState(Qt.CheckState.Checked if n % 3 == 0 else Qt.CheckState.Unchecked)
        rows[1]["deferred"] = n % 3 == 0
        self.w.approve_scan_button.click()
        context = {**self.expected[f"M{n:04}"]["snapshot"], "expedition": self.expedition}
        self.record("Ritual", {"page": 1, "items": rows, "raw_text": raw,
                               "tribute_available": 20000 + n, "rerolls_remaining": n % 4}, context, reference="Page 1")
        self.expected[f"M{n:04}"]["ritual"] = rows
        pages = logger.ritual_pages_for_map(f"M{n:04}")
        self.check(len(pages), 1, "ritual:page-count")
        self.check(pages[0]["items"], rows, "ritual:persisted-rewards")
        self.coverage["Ritual quantities/cost/deferred/counters"] += 1
        automatic_rows = [{"category": "Omen", "name": "Omen of Whittling", "quantity": 2,
                           "tribute": 500 + n, "source": f"Clear Ritual offer {n}", "deferred": n % 2 == 0}]
        automatic_raw = f"Clear synthetic Ritual evidence {n}"
        clear = {"items": [{**automatic_rows[0], "score": .99, "name_match": 1}],
                 "raw_text": automatic_raw, "tribute_available": 30000 + n, "rerolls_remaining": n % 4,
                 "reward_count": 1, "unmatched": []}
        self.w._ritual_read(clear, live=True)
        self.record("Ritual", {"page": 2, "items": automatic_rows, "raw_text": automatic_raw,
                    "tribute_available": 30000 + n, "rerolls_remaining": n % 4}, context, reference="Page 2")
        self.check(self.w.pending_review_kind, None, "auto:ritual-saved-review-cleared")
        pages = logger.ritual_pages_for_map(f"M{n:04}")
        self.check(len(pages), 2, "ritual:auto-and-manual-page-history")
        self.check(pages[0]["items"], rows, "ritual:auto-keeps-earlier-page")
        self.check(pages[1]["items"], automatic_rows, "ritual:automatic-offer-fields-exact")
        self.expected[f"M{n:04}"]["ritual"] += automatic_rows
        self.coverage["clear live Ritual automatic commit"] += 1
        self.verify_totals()

    def remnant(self, n, seed=False):
        """Deliver opened or seed results and verify confidence holds and resolved recipe commits."""
        context = logger.scan_context()
        if seed:
            reading = {"sockets": 5, "seed_slot": "P3", "seed_rune": "Power", "family": "Family 1",
                       "candidates": [1], "can_commit": True, "rewards": ["Divine Orb", "Divine Orb x2"]}
            result = {"mode": "seed", "remnants": [reading], "status": "Synthetic seed reading", **context}
            expected = OPENED[1:]
        else:
            result = {"mode": "opened", "can_use": True, "family": "Family 1", "candidates": [1],
                      "sockets": 6, "recipe_sockets": 6, "socket_source": "opened icons", "list_complete": False,
                      "first_recipe": "Divine Orb x3", "next_recipe": "Divine Orb x2",
                      "first_line_gap": 40, "status": "Review the opened rewards before logging.",
                      "opened_recipes": [{"recipe": row["recipe"], "raw": row["recipe"], "ocr_score": .99, "match_score": 1}
                                         for row in OPENED[:2]], **context}
            expected = OPENED
        mode = "seed" if seed else "opened"
        if n % 10 == 0:
            held = copy.deepcopy(result)
            if seed:
                held["remnants"][0]["can_commit"] = False
            else:
                held["opened_recipes"][0]["ocr_score"] = .2
            before = len(self.commit_log)
            self.w._scan_done(mode, held, self.raw.getvalue())
            self.check(logger.get_state()["scan_commit_count"], before,
                       "auto:unclear-seed-held" if seed else "auto:unclear-opened-held")
            self.check(self.w.pending_review_kind, "seed" if seed else "remnant", "auto:remnant-review-required")
            self.w.reject_scan_button.click()
            self.coverage["unclear seed held with auto-all enabled" if seed else
                          "unclear opened remnant held with auto-all enabled"] += 1
        self.w._scan_done(mode, result, self.raw.getvalue())
        details = {"recipes": expected}
        if seed:
            details["visible_seed"] = {"sockets": 5, "slot": "P3", "rune": "Power", "scan_index": 1, "mode": "seed"}
        self.record("Remnant", details, {**self.expected[f"M{n:04}"]["snapshot"], "expedition": self.expedition})
        self.expected[f"M{n:04}"]["recipes"] += copy.deepcopy(expected)
        self.coverage["seed recipe family order" if seed else "opened recipe resolution"] += 1
        self.coverage["clear live seed automatic commit" if seed else "clear live opened remnant automatic commit"] += 1
        self.check(self.w.first_recipe.text(), "", "review:remnant-draft-cleared")
        self.check(self.w.pending_review_kind, None, "review:remnant-pending-cleared")

    def propagation(self, n):
        """Exercise automatic and reviewed chain parts, corrections and completion with ledger checks."""
        mid = f"M{n:04}"
        eid = mid + "-E01"
        steps = [{"step": 1, "rune1": "Death", "rune2": "Power"},
                 {"step": 2, "rune1": "Opulent", "rune2": ""}]
        context = {**self.expected[mid]["snapshot"], "expedition": self.expedition}
        if n % 3:
            self.w._propagation_read({"mode": "propagation", "can_use": True, "runes": ["Death", "Power"],
                "positions": [1, 2], "selected_recipe": "Divine Orb x2", "status": "Pair scan", **logger.scan_context()}, self.raw.getvalue())
            self.record("Propagation", {"runes": ["Death", "Power"], "recipe": "Divine Orb x2", "detonated": 1},
                        context, reference="Divine Orb x2", batch_remaining=1)
            self.record("Chain", {"steps": steps[:1], "detonated": 1}, context, reference="Steps 1–1")
            self.coverage["confident propagation automatically saves chain part"] += 1
        else:
            clear_candidate = n % 6 == 0
            self.w._propagation_read({"mode": "propagation", "can_use": False, "runes": [],
                "status": "Choose recipe", "choices": [
                    {"selected_recipe": "Wrong recipe", "runes": ["Rage", "Time"], "can_use": True},
                    {"selected_recipe": "Divine Orb x2", "runes": ["Death", "Power"] if clear_candidate else [],
                     "can_use": clear_candidate}], **logger.scan_context()}, self.raw.getvalue())
            self.w.propagation_recipe_table.cellWidget(0, 2).findChildren(QPushButton)[1].click()
            self.w.propagation_recipe_table.setCurrentCell(1, 0)
            if clear_candidate:
                self.w.propagation_recipe_table.cellWidget(1, 2).findChildren(QPushButton)[0].click()
                self.coverage["held recipe row approval saves directly"] += 1
            else:
                for field, rune in zip(self.w._propagation_row_inputs[1], ("Death", "Power")):
                    self.choose(field, rune)
                self.w.propagation_recipe_table.cellWidget(1, 2).findChild(
                    QPushButton, "approvePropagationRecipe").click()
                self.coverage["held recipe dropdown corrections save directly"] += 1
            self.record("Propagation", {"runes": ["Death", "Power"], "recipe": "Divine Orb x2", "detonated": 1},
                        context, reference="Divine Orb x2", batch_remaining=1)
            self.check(logger.get_state()["chain"], steps[:1], "propagation:manual-part-saved-directly")
            self.check(self.w.chain_review_table.rowCount(), 0, "propagation:no-manual-paired-draft")
            self.check(self.w.review_complete_chain_button.isEnabled(), True,
                       "propagation:reviewed-part-ready-for-completion")
            self.record("Chain", {"steps": steps[:1], "detonated": 1}, context, reference="Steps 1–1")
            self.coverage["propagation deny and row approval"] += 1
        self.check(logger.get_state()["current_expedition_id"], eid, "propagation:first-part-stays-in-expedition")
        self.check(logger.get_state()["chain"], steps[:1], "propagation:first-saved-paired-part")
        self.check(self.w.chain_review_table.rowCount(), 0, "propagation:saved-part-clears-draft")
        self.w._propagation_read({"mode": "propagation", "can_use": True, "runes": ["Opulent"], "positions": [1],
            "selected_recipe": "Greater Regal Orb x3", "status": "Second chain part", **logger.scan_context()}, self.raw.getvalue())
        self.record("Propagation", {"runes": ["Opulent"], "recipe": "Greater Regal Orb x3", "detonated": 2},
                    context, reference="Greater Regal Orb x3", batch_remaining=1)
        self.record("Chain", {"steps": steps[1:], "detonated": 2}, context, reference="Steps 2–2")
        self.coverage["confident propagation automatically saves chain part"] += 1
        self.check(logger.get_state()["current_expedition_id"], eid, "propagation:repeated-auto-save-same-expedition")
        self.check(logger.get_state()["chain"], steps, "propagation:saved-part-order")
        self.check(logger.get_state()["detonated"], 2, "propagation:pair-counts-one")
        self.check(self.w.expedition_chain_table.rowCount(), 2, "propagation:saved-parts-visible-on-expedition")
        self.check([field.text() for field in self.w.rune_inputs], [""] * len(self.w.rune_inputs),
                   "propagation:automatic-save-clears-draft")
        if n % 10 == 0:
            with logger._connect() as db:
                identities = [tuple(row) for row in db.execute(
                    "SELECT position,chain_step,map_id,expedition_id FROM new_export "
                    "WHERE expedition_id=? AND chain_step IS NOT NULL AND chain_step!='' ORDER BY chain_step", (eid,))]
            self.w.tabs.setCurrentIndex(1)
            self.choose(self.w.expedition_chain_table.cellWidget(0, 2), "Life")
            self.check(self.w.expedition_complete_chain_button.isEnabled(), False,
                       "propagation:save-corrections-before-completing")
            self.w.expedition_save_chain_button.click()
            self.record("Chain correction", {"detonated": 2, "changes": [
                {"step": 1, "previous_rune1": "Death", "previous_rune2": "Power",
                 "rune1": "Death", "rune2": "Life"}]}, context, reference="Corrected 1 chain parts")
            steps[0]["rune2"] = "Life"
            self.check(logger.get_state()["chain"], steps, "propagation:dropdown-corrects-same-parts")
            self.check(logger.get_state()["detonated"], 2, "propagation:correction-keeps-detonated-count")
            self.check(logger.get_state()["current_expedition_id"], eid, "propagation:correction-keeps-expedition")
            with logger._connect() as db:
                corrected_identities = [tuple(row) for row in db.execute(
                    "SELECT position,chain_step,map_id,expedition_id FROM new_export "
                    "WHERE expedition_id=? AND chain_step IS NOT NULL AND chain_step!='' ORDER BY chain_step", (eid,))]
                first_part = next(entry for entry in self.commit_log if entry["kind"] == "Chain" and entry["map_id"] == mid)
                original = json.loads(db.execute("SELECT details_json FROM commits WHERE number=?",
                                                (first_part["number"],)).fetchone()[0])
            self.check(corrected_identities, identities, "propagation:correction-keeps-row-identities")
            self.check(original["steps"], first_part["details"]["steps"],
                       "propagation:correction-preserves-original-audit")
            self.coverage["saved rune dropdown corrections before completion"] += 1
        self.w.tabs.setCurrentIndex(0 if n % 2 else 1)
        completion = self.w.review_complete_chain_button if n % 2 else self.w.expedition_complete_chain_button
        self.check(completion.isEnabled(), True, "propagation:complete-chain-ready")
        completion.click()
        self.record("Chain completion", {"steps": steps, "detonated": 2,
                    "next_expedition": 2, "next_expedition_id": mid + "-E02"},
                    context, reference="Completed 2 chain parts")
        self.expected[mid]["chain"] = copy.deepcopy(steps)
        self.expedition += 1
        self.check(logger.get_state()["current_expedition_id"], mid + "-E02", "propagation:next-expedition")
        self.check(logger.get_state()["chain"], [], "propagation:next-chain-empty")
        self.check(self.w.expedition_chain_table.rowCount(), 0, "propagation:completion-clears-current-display")
        self.check(self.w.chain_list.count(), 0, "propagation:completion-clears-chain-list")
        self.check([field.text() for field in self.w.rune_inputs], [""] * len(self.w.rune_inputs), "propagation:chain-draft-cleared")
        # Use the normal expedition selector to inspect the retained final
        # chain. Completed parts are readable but their structure is locked.
        self.choose(self.w.expedition, 1)
        self.check(logger.get_state()["chain"], steps, "propagation:completed-history-retained")
        self.check(logger.get_state()["chain_completed"], True, "propagation:completed-chain-locked")
        self.check(all(not self.w.expedition_chain_table.cellWidget(row, column).isEnabled()
                       for row in range(2) for column in (1, 2)), True,
                   "propagation:completed-rune-dropdowns-locked")
        self.check(self.w.expedition_complete_chain_button.isEnabled(), False,
                   "propagation:cannot-complete-chain-twice")
        self.choose(self.w.expedition, 2)
        self.check(logger.get_state()["chain"], [], "propagation:return-to-next-empty-chain")
        self.coverage["complete chain in Review" if n % 2 else "complete chain in Expedition"] += 1
        self.coverage["paired chain parts and order"] += 1

    @staticmethod
    def map_records(db, mid):
        """Collect map-owned history and chain receipt rows for immutable-history fingerprints."""
        records = []
        for table in ("maps", "map_unique_kills", "currency_snapshots", "ritual_pages", "commits", "new_export", "expeditions"):
            records += [dict(row) for row in db.execute(f"SELECT * FROM {table} WHERE map_id=? ORDER BY rowid", (mid,))]
        for table in ("chain_completions", "chain_append_receipts"):
            records += [dict(row) for row in db.execute(
                f"SELECT * FROM {table} WHERE expedition_id LIKE ? ORDER BY rowid", (mid + "-E%",))]
        return records

    def freeze_old(self):
        """Compare every completed map's current history fingerprint with its saved fingerprint."""
        with logger._connect() as db:
            for mid, fingerprint in self.frozen.items():
                record = self.map_records(db, mid)
                self.check(hashlib.sha256(json.dumps(record, sort_keys=True).encode()).hexdigest(), fingerprint, "freeze:old-map-and-history")

    def finish(self, n):
        """Finish the map through the UI and freeze its verified counts, setup and history."""
        expected = self.expected[f"M{n:04}"]
        for field, val in zip((self.w.normal, self.w.magic, self.w.rare, self.w.unique), expected["kills"]):
            field.setText(str(val))
        if n % 10 == 0:
            self.w.rune_inputs[0].setText("Death")
            self.coverage["new-map discards unsaved chain draft"] += 1
        self.button("+ New map")
        self.record("Map totals", {"kills": expected["kills"][:3], "unique_kills": expected["kills"][3], "detonated": None},
                    expected["snapshot"], reference="")
        self.check(logger.get_state()["current_map_id"], f"M{n + 1:04}", "map:new-id")
        self.check([f.text() for f in (self.w.normal, self.w.magic, self.w.rare, self.w.unique)], [""] * 4, "map:kill-fields-cleared")
        self.check([f.text() for f in self.w.rune_inputs], [""] * len(self.w.rune_inputs), "map:chain-cleared")
        with logger._connect() as db:
            row = db.execute("SELECT * FROM maps WHERE map_id=?", (f"M{n:04}",)).fetchone()
            snapshot = json.loads(row["snapshot_json"])
            self.check(json.loads(row["kills_json"]), expected["kills"][:3], "kills:finished-map")
            unique = db.execute("SELECT unique_kills FROM map_unique_kills WHERE map_id=?", (f"M{n:04}",)).fetchone()[0]
            self.check(unique, expected["kills"][3], "kills:unique-finished-map")
            for key, val in expected["snapshot"].items():
                self.check(snapshot.get(key), val, "map-snapshot:" + key)
            data = self.map_records(db, f"M{n:04}")
        self.frozen[f"M{n:04}"] = hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest()
        self.completed = n
        self.expedition = 1

    def export(self, n, folder=False):
        """Save CSV and XLSX through UI actions and compare reopened values with the ledger."""
        destination = self.output / f"share-{n:03}"
        destination.mkdir(exist_ok=True)
        if folder:
            with patch.object(QFileDialog, "getExistingDirectory", return_value=str(destination)):
                self.button("Choose folder…", 5)
            self.button("Save CSV to folder", 5)
            self.wait_tasks()
            self.button("Save XLSX to folder", 5)
            self.wait_tasks()
            csv_path = next(p for p in destination.glob("*.csv") if not p.stem.endswith(("_Atlas", "_Scan_History")))
            xlsx_path = next(destination.glob("*.xlsx"))
            # A second folder export must preserve the first trio, even with same timestamp.
            original = {p.name: p.read_bytes() for p in destination.iterdir()}
            self.button("Save CSV to folder", 5)
            self.wait_tasks()
            for name, data in original.items():
                self.check((destination / name).read_bytes() == data, True, "exports:folder-no-overwrite")
            self.check(len(list(destination.glob("*.csv"))), 6, "exports:second-csv-trio")
            self.coverage["folder selection and collision preservation"] += 1
        else:
            csv_path = destination / "Session.csv"
            xlsx_path = destination / "Session.xlsx"
            for kind, path in (("CSV", csv_path), ("XLSX", xlsx_path)):
                action = next(a for a in self.w.findChildren(QAction) if a.text() == f"Save Export {kind} as…")
                with patch.object(QFileDialog, "getSaveFileName", return_value=(str(path), "")):
                    action.trigger()
                    self.wait_tasks()
        paths = [csv_path, csv_path.with_name(csv_path.stem + "_Atlas.csv"),
                 csv_path.with_name(csv_path.stem + "_Scan_History.csv")]
        csv_sheets = []
        for path in paths:
            with path.open(encoding="utf-8-sig", newline="") as f:
                csv_sheets.append(list(csv.DictReader(f)))
        workbook = load_workbook(xlsx_path, read_only=False, data_only=False)
        self.check(workbook.sheetnames[:3], ["Export", "Atlas Character Settings", "Scan History"], "exports:xlsx-sheet-mapping")
        for index, rows in enumerate(csv_sheets):
            sheet = workbook.worksheets[index]
            data = list(sheet.values)
            headers = list(data[0])
            self.check(len(set(headers)), len(headers), "exports:xlsx-no-duplicate-headers")
            self.check(len(data) - 1, len(rows), "exports:csv-xlsx-row-count")
            # CSV/XLSX share semantic values. Their totals are validated independently below.
            for actual, expected_row in zip(data[1:], rows):
                normalized = {h: "" if v is None else str(v) for h, v in zip(headers, actual)}
                for key, val in expected_row.items():
                    numeric = normalized.get(key, "")
                    try:
                        equal = float(numeric) == float(val) if numeric and val else numeric == val
                    except ValueError:
                        equal = numeric == val
                    self.check(equal, True, "exports:csv-xlsx-value")
        primary, atlas_rows, history = csv_sheets
        grouped = defaultdict(list)
        for row in primary:
            grouped[row["Map ID"]].append(row)
        self.check(set(grouped), set(self.expected) | {f"M{n + 1:04}"}, "exports:all-map-ids")
        atlas_ids = {row["Atlas Setup ID"] for row in atlas_rows}
        for mid, expected in self.expected.items():
            rows = grouped[mid]
            self.check(len(rows), len(expected["recipes"]) + 2, "exports:primary-recipe-chain-count")
            self.check(sum(row["Type"] == "Chain" for row in rows), 2, "exports:primary-chain-count")
            chain_rows = [row for row in rows if row["Type"] == "Chain"]
            self.check([{"step": int(row["Chain Step #"]),
                         "rune1": row["Propagation Rune 1"],
                         "rune2": row["Propagation Rune 2"]} for row in chain_rows],
                       expected["chain"], "exports:final-chain-order-and-corrections")
            self.check([row["Expedition ID"] for row in chain_rows], [mid + "-E01"] * 2,
                       "exports:chain-parts-retain-expedition")
            first = rows[0]
            phases = expected["inventories"]
            for name in self.labels | {"Chaos Orb", "Divine Orb", "Exalted Orb"}:
                quantity = max(0, phases["end"].get(name, 0) - phases.get("start", {}).get(name, 0))
                self.check(first.get(name), str(quantity), "exports:latest-positive-item-count")
                self.check([r.get(name) for r in rows[1:]], [""] * (len(rows) - 1), "exports:counts-once-per-map")
            for column, qty in zip(("Normal Kills (Map)", "Magic Kills (Map)", "Rare Kills (Map)", "Unique Kills (Map)"), expected["kills"]):
                self.check(first.get(column), str(qty), "exports:kill-count")
            self.check(first.get("Expedition 1 Detonated"), "2", "exports:paired-detonated")
            self.check(first.get("Start Baseline"), "Scanned" if "start" in phases else "Assumed empty", "exports:baseline")
            self.check(first.get("Atlas Setup ID"), expected["snapshot"]["atlas_setup_id"], "exports:atlas-map-link")
            self.check(first.get("Atlas Setup ID") in atlas_ids, True, "exports:atlas-link-resolves")
            self.check(first.get("Gear Item Rarity %"), str(expected["snapshot"]["gear_item_rarity"]), "exports:gear-rarity")
        history_commits = {int(row["Scan Commit #"]) for row in history if row.get("Scan Commit #")}
        self.check(history_commits, set(range(1, len(self.commit_log) + 1)), "exports:complete-audit-history")
        self.check(any(row.get("Type") == "Currency" for row in primary), False, "exports:primary-no-inventory-duplicates")
        self.check(any(row.get("Type") == "Ritual" for row in primary), False, "exports:primary-no-ritual-duplicates")
        for label in self.rejected_labels:
            self.check(label in primary[0], False, "exports:rejected-label-excluded")
        # Resolve actual workbook hyperlinks into the compact Atlas worksheet.
        atlas_col = [c.value for c in workbook.worksheets[0][1]].index("Atlas Setup ID") + 1
        linked = [workbook.worksheets[0].cell(r, atlas_col) for r in range(2, workbook.worksheets[0].max_row + 1)]
        self.check(all(cell.hyperlink for cell in linked if cell.value), True, "exports:xlsx-atlas-hyperlinks")
        workbook.close()
        self.exports.append({"maps": n, "csv": str(csv_path), "xlsx": str(xlsx_path),
                             "primary_rows": len(primary), "history_rows": len(history), "atlas_rows": len(atlas_rows)})
        self.coverage["real CSV and openpyxl workbook reopen"] += 1

    def restart(self):
        """Reopen the logger and verify retained map context, learned references and session totals."""
        self.w.close()
        self.w.pool.shutdown(wait=True, cancel_futures=True)
        self.app.processEvents()
        logger._READY = False
        logger.initialize()
        self.w = LoggerWindow()
        self.w._poll.stop()
        self.w.resize(1800, 1200)
        self.w.show()
        self.app.processEvents()
        self.check(logger.get_state()["current_map_id"], f"M{self.current + 1:04}", "restart:current-map")
        self.check(self.labels <= set(logger.inventory_names()), True, "restart:learned-labels")
        self.check(any(r["name"] in self.labels for r in logger.review_icons()), True, "restart:learned-artwork")
        self.check(self.w.auto_all_checkbox.isChecked(), True, "restart:auto-all-ui-preference")
        self.verify_totals()
        self.freeze_old()
        self.coverage["application restart persistence"] += 1

    def test_one_hundred_maps_saved_and_shared(self):
        """Run 100 varied maps and validate scans, chains, exports, backup and restart persistence."""
        self.w.auto_all_checkbox.setChecked(True)
        self.app.processEvents()
        self.check({key: logger.get_state()["settings"][key] for key in
                    ("ocr_auto_commit", "auto_commit", "tablet_auto_commit")},
                   {"ocr_auto_commit": True, "auto_commit": True, "tablet_auto_commit": True},
                   "auto:all-activities-enabled-through-ui")
        self.coverage["auto-all enabled through Options checkbox"] += 1
        for n in range(1, 101):
            self.current = n
            self.freeze_old()
            self.configure(n)
            if n % 4:
                start = {} if n % 11 == 0 else {"Chaos Orb": n * 3, "Divine Orb": n % 4 + 3}
                self.inventory("start", start, automatic=bool(start))
            end = {"Chaos Orb": n * 3 + n % 9, "Divine Orb": n % 4,
                   "Exalted Orb": n % 13}
            custom = f"User labelled token {n // 10}" if n % 10 == 1 else None
            rejected = f"Rejected token {n}" if custom else None
            self.inventory("end", end, custom=custom, rejected=rejected, unnamed=bool(custom), live=bool(custom))
            if n % 5 == 0:
                self.reject_inventory()
                self.inventory("end", {"Chaos Orb": n * 3 + 1, "Exalted Orb": n % 7})
                self.coverage["latest End correction"] += 1
            # Repeat the actual End capture with clear stacks. It replaces the
            # current map snapshot while retaining the earlier review audit.
            self.inventory("end", self.expected[f"M{n:04}"]["inventories"]["end"], automatic=True)
            self.ritual(n)
            self.remnant(n)
            self.remnant(n, seed=True)
            self.propagation(n)
            self.finish(n)
            if n % 25 == 0:
                self.export(n, folder=n in (25, 75))
                if n != 100:
                    self.restart()
            if n % 10 == 0:
                print(f"100-map GUI progress: {n}/100, {len(self.commit_log)} commits, {len(self.failures)} mismatches", flush=True)
        for activity in ("Atlas settings", "clear live waystone automatic commit",
                         "clear live tablet automatic commit", "clear live currency automatic commit",
                         "clear live Ritual automatic commit", "clear live opened remnant automatic commit",
                         "clear live seed automatic commit", "paired chain parts and order", "Map totals"):
            self.check(self.coverage[activity] >= 100, True, "coverage:at-least-100:" + activity)
        self.check(self.coverage["confident propagation automatically saves chain part"] >= 100,
                   True, "coverage:at-least-100:automatic-propagation")
        for activity in ("unclear waystone held with auto-all enabled", "unclear tablet held with auto-all enabled",
                         "unclear currency held with auto-all enabled", "unclear Ritual held with auto-all enabled",
                         "unclear opened remnant held with auto-all enabled", "unclear seed held with auto-all enabled"):
            self.check(self.coverage[activity] >= 10, True, "coverage:uncertain-reads-stay-held:" + activity)
        # Backup is exported through the File action and opened independently as SQLite.
        backup = self.output / "Session_Backup.sqlite3"
        action = next(a for a in self.w.findChildren(QAction) if a.text() == "Save database backup as…")
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(backup), "")):
            action.trigger()
            self.wait_tasks()
        with sqlite3.connect(backup) as db:
            self.check(db.execute("SELECT count(*) FROM commits").fetchone()[0], len(self.commit_log), "backup:commits")
            self.check(db.execute("SELECT count(*) FROM maps").fetchone()[0], 101, "backup:maps")
            self.check(db.execute("PRAGMA integrity_check").fetchone()[0], "ok", "backup:integrity")
        self.coverage["menu backup export independently reopened"] += 1
        self.restart()
        self.assertEqual(self.failures, [], f"{len(self.failures)} mismatches; see {self.output / 'metrics.json'}")


if __name__ == "__main__":
    unittest.main()

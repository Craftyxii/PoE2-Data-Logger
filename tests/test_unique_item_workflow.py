"""Keep full-item recognition proposals in existing review, learning and export paths.

These tests deliver declared recognition results to real Qt callbacks; textured
fixtures test evidence ownership and storage, not unique-art recognition accuracy.
"""

import csv
import io
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np
from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import QSize, Qt
from PySide6.QtWidgets import QApplication, QPushButton

from PoE2_Data_Logger.core import logger_store as logger, review_learning, store, workbook_export
from PoE2_Data_Logger.ui.native_desktop import LoggerWindow


class UniqueItemWorkflowTests(unittest.TestCase):
    """Keep preloaded Item proposals separate from Currency catalog registration."""

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="poe2-unique-workflow-")
        self.previous = store.DATA_DIR
        store.DATA_DIR = Path(self.temporary.name)
        logger._READY = False
        logger.initialize()
        self.window = LoggerWindow()
        self.window._poll.stop()
        self.jobs = []
        self.window._submit = lambda label, work, done=None: self.jobs.append(done)
        self.grid = Image.new("RGB", (648, 270), (8, 8, 22))
        self.grid.info["poe2_inventory_aligned"] = True
        pixels = np.random.default_rng(614).integers(25, 235, (162, 108, 3), dtype=np.uint8)
        self.gear = Image.fromarray(pixels)
        self.grid.paste(self.gear, (54, 54))
        self.page = Image.new("RGB", (648, 600), (8, 8, 22))
        self.page.paste(self.gear, (50, 50))

    def tearDown(self):
        self.window.close()
        self.window.pool.shutdown(wait=True, cancel_futures=True)
        self.app.processEvents()
        store.DATA_DIR = self.previous
        logger._READY = False
        self.temporary.cleanup()

    def proposal(self, name="Catalog Unique Armour", **changes):
        return {"slot": 14, "name": name, "quantity": 1, "score": .995,
                "category": "Item", "catalog_unique": True, "catalog_revision": "fixture-revision",
                "grid_slots": [14, 15, 26, 27, 38, 39], "box": [54, 54, 162, 216],
                "count_needs_review": False, "name_needs_review": False, **changes}

    def inventory(self, item):
        self.window._inventory_captured(self.grid, live=False)
        self.jobs[-1]({"items": [item], "unknown": []})

    def approve_currency_row(self, row=0):
        controls = self.window.inventory_table.cellWidget(row, 3)
        controls.findChild(QPushButton, "approveCurrency").click()
        self.assertEqual(controls.property("reviewStatus"), "approved")

    def history(self):
        return list(csv.DictReader(io.StringIO(logger.export_all_csv().decode("utf-8-sig"))))

    def test_currency_unique_registers_only_on_commit_as_item_and_keeps_full_preview(self):
        self.inventory(self.proposal())
        self.assertNotIn("Catalog Unique Armour", logger.inventory_names())
        self.assertEqual(logger.get_state()["scan_commit_count"], 0)
        icon = self.window.inventory_table.item(0, 0).icon()
        size = icon.actualSize(QSize(48, 48))
        self.assertEqual((size.width(), size.height()), (32, 48))
        self.window.approve_review()
        self.assertIn("Catalog Unique Armour", logger.item_names())
        self.assertNotIn("Catalog Unique Armour", logger.currency_names())
        self.assertEqual(logger.currency_for_map("M0001")["start"], {"Catalog Unique Armour": 1})
        self.assertEqual(logger.review_icons(), [])  # Unchanged catalog art is not a correction.
        rows = self.history()
        self.assertTrue(any(row.get("Item: Catalog Unique Armour") == "1" for row in rows))
        self.assertFalse(any("Currency: Catalog Unique Armour" in row for row in rows))
        book = load_workbook(io.BytesIO(workbook_export.export_xlsx()), read_only=True, data_only=True)
        try:
            history = list(book["Scan History"].values)
            column = history[0].index("Item: Catalog Unique Armour")
            self.assertIn(1, [row[column] for row in history[1:]])
        finally:
            book.close()
        logger.undo_last_action()
        self.assertNotIn("Catalog Unique Armour", logger.item_names())
        self.assertEqual(logger.get_state()["scan_commit_count"], 0)

    def test_currency_shared_art_choice_keeps_one_quantity_and_learns_complete_item(self):
        self.inventory(self.proposal("", candidate="Candidate Armour A / Candidate Armour B",
                                     name_needs_review=True))
        selector = self.window.inventory_table.cellWidget(0, 1)
        self.assertEqual(selector.count(), 2)
        self.assertEqual(selector.currentText(), "")
        self.assertEqual(self.window.inventory_table.item(0, 2).text(), "1")
        controls = self.window.inventory_table.cellWidget(0, 3)
        self.assertTrue(controls.property("reviewNameRequired"))
        self.assertFalse(controls.property("reviewCountRequired"))
        selector.setCurrentIndex(1)
        self.approve_currency_row()
        self.window.approve_review()
        self.assertIn("Candidate Armour B", logger.item_names())
        examples = logger.review_icons()
        self.assertEqual([(row["name"], row["kind"], row["columns"], row["rows"])
                          for row in examples], [("Candidate Armour B", "item", 2, 3)])
        with Image.open(io.BytesIO(examples[0]["image"])) as saved:
            self.assertEqual(saved.size, (192, 288))
            expected = self.gear.resize((192, 288), Image.Resampling.LANCZOS)
            self.assertEqual(saved.convert("RGB").tobytes(), expected.tobytes())
        self.assertEqual(logger.undo_last_action()["action"], "Currency approval")
        self.assertEqual(logger.review_icons(), [])
        self.assertNotIn("Candidate Armour B", logger.item_names())
        self.assertEqual(logger.currency_for_map("M0001")["start"], {})

    def test_currency_pending_unique_is_not_registered_or_learned(self):
        self.inventory(self.proposal("", candidate="Unapproved Armour A / Unapproved Armour B",
                                     name_needs_review=True))
        self.window.approve_review()
        self.assertEqual(logger.currency_for_map("M0001")["start"], {})
        self.assertNotIn("Unapproved Armour A", logger.item_names())
        self.assertNotIn("Unapproved Armour B", logger.item_names())
        self.assertEqual(logger.review_icons(), [])

    def test_currency_name_correction_to_existing_currency_keeps_canonical_kind(self):
        self.inventory(self.proposal())
        self.window.inventory_table.item(0, 1).setText("Chaos Orb")
        self.approve_currency_row()
        self.window.approve_review()
        self.assertEqual(logger.currency_for_map("M0001")["start"], {"Chaos Orb": 1})
        self.assertNotIn("Chaos Orb", logger.item_names())
        self.assertEqual(logger.review_icons(), [])

    def test_backend_duplicate_item_metadata_is_order_independent_and_undoable(self):
        for items in ([{"name": "New Unique Ring", "quantity": 1},
                       {"name": "New Unique Ring", "quantity": 1, "category": "Item"}],
                      [{"name": "New Unique Ring", "quantity": 1, "category": "Item"},
                       {"name": "New Unique Ring", "quantity": 1}]):
            with self.subTest(first_category=items[0].get("category")):
                logger.save_currency_snapshot("start", items, register_names=True)
                self.assertIn("New Unique Ring", logger.item_names())
                self.assertNotIn("New Unique Ring", logger.currency_names())
                self.assertEqual(logger.currency_for_map("M0001")["start"]["New Unique Ring"], 2)
                logger.undo_last_action()
                self.assertNotIn("New Unique Ring", logger.item_names())

    def test_backend_preserves_existing_currency_and_legacy_default_registration(self):
        logger.save_currency_snapshot("start", [
            {"name": "Chaos Orb", "quantity": 2, "category": "Item"},
            {"name": "Legacy Manual Token", "quantity": 3, "category": "Omen"}], register_names=True)
        self.assertNotIn("Chaos Orb", logger.item_names())
        self.assertIn("Legacy Manual Token", logger.currency_names())
        self.assertNotIn("Legacy Manual Token", logger.ritual_names())

    def test_inventory_learning_rejects_invalid_or_unrelated_full_item_regions(self):
        original = self.proposal("")
        for changes in ({"box": [0, 0, 108, 162]}, {"slot": True},
                        {"grid_slots": [14, 16, 26, 28]}, {"grid_slots": [61]},
                        {"box": [54, 54, float("nan"), 216]}):
            with self.subTest(changes=changes):
                self.assertIsNone(review_learning.make_example(
                    self.grid, {**original, **changes}, "Corrected Armour", "Item", "currency"))
        self.assertIsNone(review_learning.make_example(
            self.grid, original, "Chaos Orb", "Currency", "currency"))

    def test_ritual_shared_unique_uses_existing_dropdown_and_exports_item_quantity(self):
        self.window._ritual_captured(self.page, live=False)
        self.jobs[-1]({"items": [{"name": "", "category": "Item", "category_verified": True,
            "quantity": 1, "tribute": 10, "deferred": False, "needs_review": True,
            "name_needs_review": True, "count_needs_review": False, "unresolved": True,
            "candidate_names": ["Ritual Unique A", "Ritual Unique B"],
            "catalog_unique": True, "catalog_revision": "fixture-revision",
            "grid_slots": [1, 2, 13, 14, 25, 26], "box": [50, 50, 158, 212],
            "source": "unique item catalog: fixture-revision"}], "unmatched": [],
            "raw_text": "", "tribute_available": 200, "rerolls_remaining": 2})
        selector = self.window.ritual_table.cellWidget(0, 1)
        self.assertEqual(selector.count(), 2)
        self.assertEqual(selector.currentText(), "")
        selector.setCurrentIndex(0)
        # Choosing a candidate identifies the reward; its explicit row decision
        # accepts it for the later atomic page save and artwork learning.
        controls = self.window.ritual_table.cellWidget(0, 6)
        controls.findChild(QPushButton, "approveRitual").click()
        self.assertEqual(controls.property("reviewStatus"), "approved")
        self.window.approve_review()
        self.assertIn("Ritual Unique A", logger.item_names())
        pages = logger.ritual_pages_for_map("M0001")
        self.assertEqual((pages[0]["items"][0]["name"], pages[0]["items"][0]["category"],
                          pages[0]["items"][0]["quantity"]), ("Ritual Unique A", "Item", 1))
        self.assertTrue(any(row.get("Item: Ritual Unique A") == "1" for row in self.history()))
        self.assertEqual([(ref["columns"], ref["rows"]) for ref in logger.review_icons()], [(2, 3)])
        book = load_workbook(io.BytesIO(workbook_export.export_xlsx()), read_only=True, data_only=True)
        try:
            history = list(book["Scan History"].values)
            column = history[0].index("Item: Ritual Unique A")
            self.assertIn(1, [row[column] for row in history[1:]])
        finally:
            book.close()

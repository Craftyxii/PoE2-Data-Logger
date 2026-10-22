"""Read actual grayscale Ritual pixels through review, learning and exported history."""

import csv
import io
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QComboBox

from PoE2_Data_Logger.core import logger_store as logger, store, workbook_export
from PoE2_Data_Logger.ui.native_desktop import LoggerWindow


class RitualGrayscaleWorkflowTests(unittest.TestCase):
    """Keep a manually chosen grayscale label from becoming automatic identity later."""

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="poe2-gray-ritual-workflow-")
        self.previous = store.DATA_DIR
        store.DATA_DIR = Path(self.temporary.name)
        logger._READY = False
        logger.initialize()
        self.window = LoggerWindow()
        self.window._poll.stop()
        self.window.resize(1366, 900)
        self.window.show()
        self.jobs = []
        self.window._submit_scan = lambda label, request, done, *args: self.jobs.append((request, done))
        with Image.open(Path(__file__).parent / "fixtures/ritual_grayscale/264_tribute.png") as source:
            self.image = source.convert("RGB")
        self.app.processEvents()

    def tearDown(self):
        self.window.close()
        self.window.pool.shutdown(wait=True, cancel_futures=True)
        self.app.processEvents()
        store.DATA_DIR = self.previous
        logger._READY = False
        self.temporary.cleanup()

    def read_actual_capture(self):
        """Run the production request's recognizer on its captured pixels and references."""
        self.window._ritual_captured(self.image, live=True)
        request, deliver = self.jobs[-1]
        result = request.function(*request.args, **request.kwargs)
        deliver(result)
        self.app.processEvents()
        return result

    def slot_rows(self):
        """Locate review actions by complete source footprint rather than mutable row order."""
        table = self.window.ritual_table
        return {tuple((table.item(row, 0).data(Qt.ItemDataRole.UserRole) or {})["grid_slots"]): row
                for row in range(table.rowCount())}

    def test_actual_pixels_choice_history_and_learning_keep_grayscale_siblings_ambiguous(self):
        """A saved explicit choice retains its audit data without naming later gray offers."""
        drafts = ("123", "4", "5", "6")
        kill_fields = (self.window.normal, self.window.magic, self.window.rare, self.window.unique)
        for field, value in zip(kill_fields, drafts):
            field.setText(value)
        result = self.read_actual_capture()
        self.assertEqual([item["grid_slots"] for item in result["items"]],
                         [[1], [2], [13], [14], [25], [37], [49], [61],
                          [73, 74, 85, 86], [97], [109]])
        self.assertEqual((result["tribute_available"], result["rerolls_remaining"]), (264, 13))
        self.assertEqual(logger.get_state()["scan_commit_count"], 0)
        self.assertEqual(logger.review_icons(), [])
        self.assertEqual(self.window.pending_review_kind, "ritual")
        table = self.window.ritual_table
        rows = self.slot_rows()
        self.assertEqual(table.rowCount(), 11)
        armor_row = rows[(73, 74, 85, 86)]
        self.assertEqual((table.item(armor_row, 0).text(), table.item(armor_row, 2).text()),
                         ("Item", "1"))
        self.assertTrue({(73,), (74,), (85,), (86,)}.isdisjoint(rows))
        choices = {2: "Omen of Resurgence", 49: "Omen of Amelioration"}
        alternatives = {}
        for slot, name in choices.items():
            row = rows[(slot,)]
            selector = table.cellWidget(row, 1)
            self.assertIsInstance(selector, QComboBox)
            self.assertEqual(selector.currentText(), "")
            alternatives[slot] = [selector.itemText(i) for i in range(selector.count())]
            self.assertEqual(len(alternatives[slot]), 6 if slot == 2 else 5)
            self.assertIn(name, alternatives[slot])
            self.assertEqual(table.item(row, 5).checkState(), Qt.CheckState.Unchecked)
            self.window.tabs.widget(0).ensureWidgetVisible(table)
            selector.showPopup()
            self.app.processEvents()
            index = selector.model().index(selector.findText(name), 0)
            QTest.mouseClick(selector.view().viewport(), Qt.MouseButton.LeftButton,
                             pos=selector.view().visualRect(index).center())
            self.app.processEvents()
            self.assertEqual(table.item(row, 1).text(), name)
            self.assertEqual(table.item(row, 0).text(), "Omen")
            # Explicitly confirm the visible single count, retaining its separate review gate.
            table.item(row, 2).setText("1")
        for footprint, row in rows.items():
            if footprint not in {(2,), (49,)}:
                table.item(row, 1).setText("")
        self.assertEqual(logger.review_icons(), [])
        self.assertEqual(logger.get_state()["scan_commit_count"], 0)
        QTest.mouseClick(self.window.approve_scan_button, Qt.MouseButton.LeftButton)
        self.app.processEvents()
        pages = logger.ritual_pages_for_map("M0001")
        self.assertEqual(len(pages), 1)
        self.assertEqual([(item["name"], item["category"], item["quantity"], item["deferred"])
                          for item in pages[0]["items"]],
                         [(name, "Omen", 1, False) for name in choices.values()])
        self.assertEqual({ref["name"] for ref in logger.review_icons()}, set(choices.values()))
        self.assertEqual(logger.get_state()["scan_commit_count"], 1)
        self.assertEqual(tuple(field.text() for field in kill_fields), drafts)

        history = list(csv.DictReader(io.StringIO(logger.export_all_csv().decode("utf-8-sig"))))
        saved = [row for row in history if row["Scan Commit #"] == "1"]
        self.assertEqual([row["Reward Name"] for row in saved], list(choices.values()))
        self.assertTrue(all(row["Map ID"] == "M0001" and row["Deferred"] == "False" and
                            row["Ritual Tribute Available"] == "264" and
                            row["Ritual Rerolls Remaining"] == "13" for row in saved))
        workbook = load_workbook(io.BytesIO(workbook_export.export_xlsx()), read_only=True)
        try:
            values = workbook["Scan History"].iter_rows(values_only=True)
            headers = next(values)
            exported = [dict(zip(headers, row)) for row in values]
            self.assertEqual([row["Reward Name"] for row in exported if row["Scan Commit #"] == 1],
                             list(choices.values()))
        finally:
            workbook.close()

        repeated = self.read_actual_capture()
        self.assertEqual(logger.get_state()["scan_commit_count"], 1)
        self.assertEqual(len(logger.ritual_pages_for_map("M0001")), 1)
        self.assertEqual([item["grid_slots"] for item in repeated["items"]],
                         [item["grid_slots"] for item in result["items"]])
        rows = self.slot_rows()
        for slot in choices:
            selector = table.cellWidget(rows[(slot,)], 1)
            self.assertIsInstance(selector, QComboBox)
            self.assertEqual(selector.currentText(), "")
            self.assertEqual([selector.itemText(i) for i in range(selector.count())], alternatives[slot])
            self.assertFalse(next(item for item in repeated["items"] if item["grid_slots"] == [slot])["deferred"])
        self.assertEqual(tuple(field.text() for field in kill_fields), drafts)


if __name__ == "__main__":
    unittest.main()

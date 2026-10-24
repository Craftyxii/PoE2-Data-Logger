"""Round-trip every sourced unique label through approved inventory and Ritual records.

This exercises persistence and export projection, not screenshot recognition. The
bounded API batches intentionally include duplicate approved item rows so saved
quantities and latest-scan Item accounting are checked for every catalog label.
"""

import csv
import io
from itertools import zip_longest
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from PoE2_Data_Logger.core import logger_store as logger, reference_pack, store, workbook_export
from PoE2_Data_Logger.ocr import unique_items


class UniqueCatalogExportTests(unittest.TestCase):
    """Keep all unique names as distinct Items in SQL, history, totals and XLSX."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="poe2-all-unique-exports-")
        self.previous = store.DATA_DIR
        store.DATA_DIR = Path(self.temporary.name)
        logger._READY = False
        logger.initialize()
        logger.clear_export_and_reset_ids()
        logger.start_map()

    def tearDown(self):
        store.DATA_DIR = self.previous
        logger._READY = False
        self.temporary.cleanup()

    def csv_rows(self, data):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        headers = reader.fieldnames
        self.assertEqual(len(headers), len(set(headers)))
        rows = list(reader)
        self.assertTrue(all(None not in row for row in rows))
        return headers, rows

    def catalog_names(self):
        manifest = Path(__file__).resolve().parents[1] / "PoE2_Data_Logger/unique_item_assets/manifest.json"
        source = json.loads(manifest.read_text(encoding="utf-8"))
        names = tuple(sorted({item["name"] for item in source["items"]}, key=str.casefold))
        self.assertEqual(len(source["items"]), 449)
        self.assertEqual(len(names), 441)
        self.assertEqual(unique_items.names(), names)
        return names, source["source"]["revision"]

    def test_unobserved_catalog_is_in_sql_csv_and_xlsx_without_finds_or_ui_types(self):
        names, revision = self.catalog_names()
        local_names = logger.inventory_names()
        undo_before = logger.undo_action()
        with logger._connect() as db:
            references = list(db.execute("SELECT name,source_revision FROM unique_item_catalog ORDER BY name"))
            self.assertEqual(dict(references), dict.fromkeys(names, revision))
            for table in ("commits", "currency_snapshots", "ritual_pages", "item_names", "item_icons"):
                self.assertEqual(db.execute(f"SELECT count(*) FROM {table}").fetchone()[0], 0, table)
        self.assertFalse(set(names).intersection(local_names))
        self.assertFalse(set(names).intersection(row["name"] for row in logger.session_currency_totals()["items"]))

        # The SQL export contains reference names and source identity, with no
        # synthetic quantities or scan history needed to make them available.
        sql_export = Path(self.temporary.name) / "catalog-export.sqlite3"
        sql_export.write_bytes(logger.backup_bytes())
        sql = sqlite3.connect(sql_export)
        try:
            self.assertEqual(dict(sql.execute("SELECT name,source_revision FROM unique_item_catalog")),
                             dict.fromkeys(names, revision))
            self.assertEqual(sql.execute("SELECT count(*) FROM commits").fetchone()[0], 0)
        finally:
            sql.close()

        history_headers, history = self.csv_rows(logger.export_all_csv())
        primary_headers, primary = self.csv_rows(logger.export_primary_csv())
        self.assertEqual(history, [])
        self.assertEqual(len(primary), 1)
        for name in names:
            self.assertIn("Item: " + name, history_headers)
            self.assertIn(name, primary_headers)
            self.assertEqual(primary[0][name], "")
        self.assertEqual(self.csv_rows(logger.export_currency_csv())[1], [])
        self.assertEqual(self.csv_rows(logger.export_ritual_csv())[1], [])

        book = load_workbook(io.BytesIO(workbook_export.export_xlsx()), read_only=True, data_only=True)
        try:
            for sheet_name, headers in (("Scan History", history_headers), ("Export", primary_headers)):
                rows = book[sheet_name].iter_rows(values_only=True)
                self.assertEqual(list(next(rows)), headers)
                values = list(rows)
                if sheet_name == "Scan History":
                    self.assertEqual(values, [])
                else:
                    self.assertEqual(len(values), 1)
                    self.assertTrue(all(values[0][headers.index(name)] is None for name in names))
        finally:
            book.close()

        # Reopening a profile is idempotent and does not decode recognition art.
        logger._READY = False
        with patch("PIL.Image.open", side_effect=AssertionError("Startup decoded an image")):
            logger.initialize()
        self.assertEqual(logger.inventory_names(), local_names)
        self.assertEqual(logger.undo_action(), undo_before)
        with logger._connect() as db:
            self.assertEqual(list(db.execute("SELECT name,source_revision FROM unique_item_catalog ORDER BY name")),
                             references)

    def test_catalog_keeps_existing_and_recorded_categories_and_canonical_labels(self):
        self.catalog_names()
        logger.add_currency_item("bramblejack")
        logger.add_ritual_name("HeAdHuNtEr")
        headers, _ = self.csv_rows(logger.export_all_csv())
        self.assertIn("Currency: bramblejack", headers)
        self.assertIn("Omen: HeAdHuNtEr", headers)
        self.assertNotIn("Item: Bramblejack", headers)
        self.assertNotIn("Item: Headhunter", headers)
        logger.save_currency_snapshot("end", [{"name": "Bramblejack", "quantity": 2,
                                               "category": "Currency"}], register_names=True)
        logger.save_ritual_page([{"name": "Headhunter", "quantity": 1, "category": "Omen"}],
                                register_names=True)
        for reset in (False, True):
            if reset:
                reference_pack.reset_to_defaults()
            headers, history = self.csv_rows(logger.export_all_csv())
            self.assertIn("Currency: bramblejack", headers)
            self.assertIn("Omen: HeAdHuNtEr", headers)
            self.assertNotIn("Item: Bramblejack", headers)
            self.assertNotIn("Item: Headhunter", headers)
            currency = next(row for row in history if row["Type"] == "Currency")
            offer = next(row for row in history if row["Type"] == "Ritual")
            self.assertEqual(currency["Currency: bramblejack"], "2")
            self.assertEqual((offer["Reward Type"], offer["Omen: HeAdHuNtEr"]), ("Omen", "1"))

    def test_reference_reset_import_and_undo_preserve_catalog_and_real_records(self):
        names, revision = self.catalog_names()
        logger.save_currency_snapshot("end", [{"name": "Bramblejack", "quantity": 2,
                                               "category": "Item"}], register_names=True)
        pack = Path(self.temporary.name) / "real-references.zip"
        pack.write_bytes(reference_pack.export_pack())
        with logger._connect() as db:
            saved = {table: [tuple(row) for row in db.execute(f"SELECT * FROM {table}")]
                     for table in ("commits", "currency_snapshots")}

        def unchanged():
            with logger._connect() as db:
                self.assertEqual(dict(db.execute("SELECT name,source_revision FROM unique_item_catalog")),
                                 dict.fromkeys(names, revision))
                for table, rows in saved.items():
                    self.assertEqual([tuple(row) for row in db.execute(f"SELECT * FROM {table}")], rows)
            _, history = self.csv_rows(logger.export_all_csv())
            real = next(row for row in history if row["Type"] == "Item")
            self.assertEqual(real["Item: Bramblejack"], "2")

        reference_pack.reset_to_defaults()
        unchanged()
        self.assertNotIn("Bramblejack", logger.item_names())
        reference_pack.import_pack(pack)
        unchanged()
        self.assertIn("Bramblejack", logger.item_names())
        self.assertEqual(logger.undo_last_action()["action"], "Import OCR references")
        unchanged()
        self.assertNotIn("Bramblejack", logger.item_names())
        self.assertEqual(logger.undo_last_action()["action"], "Reset OCR references")
        unchanged()
        self.assertIn("Bramblejack", logger.item_names())

    def test_every_catalog_name_round_trips_with_latest_inventory_and_ritual_offers(self):
        names, _ = self.catalog_names()
        # Export-only references do not register unapproved user data.
        self.assertFalse(set(names).intersection(logger.item_names()))
        self.assertFalse(set(names).intersection(logger.currency_names()))

        maps = {}
        for offset in range(0, len(names), 20):
            batch = names[offset:offset + 20]
            map_id = logger.get_state()["current_map_id"]
            maps[map_id] = batch

            def copies(count):
                return [{"name": name, "quantity": 1, "category": "Item"}
                        for name in batch for _ in range(count)]

            logger.save_currency_snapshot("start", copies(1), register_names=True)
            interim = logger.save_currency_snapshot("end", copies(2), register_names=True)
            self.assertEqual(interim["end"], dict.fromkeys(batch, 2))
            final = logger.save_currency_snapshot("end", copies(3), register_names=True)
            self.assertEqual(final["start"], dict.fromkeys(batch, 1))
            self.assertEqual(final["end"], dict.fromkeys(batch, 3))
            self.assertEqual(final["net"], dict.fromkeys(batch, 3))
            logger.save_ritual_page([
                {"name": name, "quantity": 1, "category": "Item", "deferred": deferred,
                 "source": "approved sourced unique catalog proposal"}
                for name in batch for deferred in (False, True)], register_names=True,
                tribute_available=1000, rerolls_remaining=2)
            logger.finish_map(0, 0, 0)
            if offset + len(batch) < len(names):
                logger.start_map()

        self.assertEqual(set(logger.item_names()), set(names))
        self.assertFalse(set(names).intersection(logger.currency_names()))
        self.assertFalse(set(names).intersection(logger.ritual_names()))
        with logger._connect() as db:
            snapshots = list(db.execute("SELECT map_id,phase,items_json FROM currency_snapshots"))
            self.assertEqual(len(snapshots), len(maps) * 2)
            for saved in snapshots:
                amount = 1 if saved["phase"] == "start" else 3
                self.assertEqual(json.loads(saved["items_json"]), dict.fromkeys(maps[saved["map_id"]], amount))
            for saved in db.execute("SELECT kind,map_id,details_json FROM commits WHERE kind IN ('Currency','Ritual')"):
                details = json.loads(saved["details_json"])
                if saved["kind"] == "Currency":
                    self.assertEqual(details["item_kinds"], dict.fromkeys(maps[saved["map_id"]], "Item"))
                else:
                    # Each name's deferred copy was logged when first offered, so it is not saved.
                    self.assertEqual(len(details["items"]), len(maps[saved["map_id"]]))
                    self.assertTrue(all(item["category"] == "Item" and item["quantity"] == 1
                                        for item in details["items"]))
            pages = list(db.execute("SELECT map_id,items_json FROM ritual_pages"))
            self.assertEqual(len(pages), len(maps))
            for saved in pages:
                offers = json.loads(saved["items_json"])
                self.assertEqual({item["name"] for item in offers}, set(maps[saved["map_id"]]))
                self.assertTrue(all(item["category"] == "Item" for item in offers))

        totals = logger.session_currency_totals()
        self.assertEqual(totals["maps_counted"], len(maps))
        self.assertEqual({row["name"]: (row["kind"], row["quantity"]) for row in totals["items"]},
                         {name: ("Item", 3) for name in names})

        _, inventory = self.csv_rows(logger.export_currency_csv())
        self.assertEqual(len(inventory), len(names))
        self.assertEqual({row["Currency"] for row in inventory}, set(names))
        for row in inventory:
            self.assertIn(row["Currency"], maps[row["Map ID"]])
            self.assertEqual([row[key] for key in ("Start Count", "End Count", "Net Change", "Session Found Quantity")],
                             ["1", "3", "3", "3"])
            self.assertEqual(row["End Currency Scan #"], "3")

        _, ritual = self.csv_rows(logger.export_ritual_csv())
        self.assertEqual(len(ritual), len(names))
        self.assertEqual({row["Name"] for row in ritual}, set(names))
        for row in ritual:
            self.assertIn(row["Name"], maps[row["Map ID"]])
            self.assertEqual((row["Type"], row["Quantity"]), ("Item", "1"))
            self.assertEqual((row["Deferred"], row["New Find Quantity"]), ("False", "1"))

        headers, history = self.csv_rows(logger.export_all_csv())
        for name in names:
            self.assertIn("Item: " + name, headers)
            self.assertNotIn("Currency: " + name, headers)
            self.assertNotIn("Omen: " + name, headers)
        found = {name: {"inventory": [], "offers": []} for name in names}
        for row in history:
            if row["Type"] == "Item":
                name = row["Item Name"]
                self.assertIn(name, maps[row["Map ID"]])
                self.assertEqual(row["Currency"], "")
                self.assertEqual(row["Item: " + name], row["Quantity"])
                found[name]["inventory"].append(row)
            elif row["Type"] == "Ritual":
                name = row["Reward Name"]
                self.assertEqual(row["Reward Type"], "Item")
                self.assertEqual(row["Item: " + name], row["New Find Quantity"])
                found[name]["offers"].append(row)
        for name, readings in found.items():
            self.assertEqual([row["Quantity"] for row in readings["inventory"]], ["1", "2", "3"], name)
            self.assertEqual([row["Final Currency Scan"] for row in readings["inventory"]],
                             ["False", "False", "True"], name)
            self.assertEqual([row["Session Found Quantity"] for row in readings["inventory"]],
                             ["0", "0", "3"], name)
            self.assertEqual([row["Item: " + name] for row in readings["offers"]], ["1"], name)

        primary_headers, primary = self.csv_rows(logger.export_primary_csv())
        self.assertEqual(len(primary), len(maps))
        for row in primary:
            for name in maps[row["Map ID"]]:
                self.assertIn(name, primary_headers)
                self.assertEqual(row[name], "3")

        book = load_workbook(io.BytesIO(workbook_export.export_xlsx()), read_only=True, data_only=True)
        try:
            for sheet_name, csv_headers, expected_rows in (("Scan History", headers, history),
                                                          ("Export", primary_headers, primary)):
                rows = book[sheet_name].iter_rows(values_only=True)
                self.assertEqual(list(next(rows)), csv_headers)
                type_index = csv_headers.index("Type")
                name_index = csv_headers.index("Item Name" if sheet_name == "Scan History" else "Map ID")
                offer_index = csv_headers.index("Reward Name") if sheet_name == "Scan History" else None
                item_indices = {name: csv_headers.index("Item: " + name if sheet_name == "Scan History" else name)
                                for name in names}
                for expected, values in zip_longest(expected_rows, rows):
                    self.assertIsNotNone(expected)
                    self.assertIsNotNone(values)
                    self.assertEqual(["" if value is None else str(value) for value in values],
                                     [expected[header] for header in csv_headers])
                    if sheet_name == "Scan History" and values[type_index] in ("Item", "Ritual"):
                        name = values[name_index] if values[type_index] == "Item" else values[offer_index]
                        self.assertIs(type(values[item_indices[name]]), int)
                    elif sheet_name == "Export":
                        for name in maps[values[name_index]]:
                            self.assertIs(type(values[item_indices[name]]), int)
                            self.assertEqual(values[item_indices[name]], 3)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
